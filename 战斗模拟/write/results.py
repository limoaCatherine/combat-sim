"""把一场结果追加到「对战模拟」的运行结果块。不覆盖已有运行。"""
from __future__ import annotations

from datetime import datetime

import openpyxl

from 战斗模拟.config import DATA_START, SHEET_TASK
from 战斗模拟.load.blocks import block_spans


def _headers(ws, start: int, end: int) -> dict[str, int]:
    out = {}
    for col in range(start, end + 1):
        name = ws.cell(3, col).value
        if name:
            out[str(name).strip()] = col
    return out


def append_results(path: str, rows: list[dict]) -> int:
    wb = openpyxl.load_workbook(path)
    try:
        ws = wb[SHEET_TASK]
        span = next(item for item in block_spans(ws) if item[0] == "运行结果")
        _, start, end = span
        headers = _headers(ws, start, end)
        id_col = headers.get("运行ID", start)
        row = DATA_START
        while ws.cell(row, id_col).value not in (None, ""):
            row += 1
        for item in rows:
            for key, value in item.items():
                col = headers.get(key)
                if col:
                    ws.cell(row, col).value = value
            row += 1
        wb.save(path)
        return row - 1
    finally:
        wb.close()


def result_row(run_id: str, task: str, scene: str, match: str, unit: dict, duration_ms: float, win: str) -> dict:
    seconds = duration_ms / 1000 if duration_ms else 0
    damage = unit.get("伤害") or 0
    return {
        "运行ID": run_id,
        "运行时间": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "任务名": task,
        "场景名": scene,
        "对阵": match,
        "阵营": unit.get("阵营"),
        "主体名": unit.get("名称"),
        "算法": "掷骰",
        "样本数": 1,
        "均值": damage,
        "标准差": 0,
        "CV": 0,
        "P05": damage,
        "P50": damage,
        "P95": damage,
        "指标名": "伤害",
    }
