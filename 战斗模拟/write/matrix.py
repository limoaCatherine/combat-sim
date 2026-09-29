# -*- coding: utf-8 -*-
"""把一场战斗的指标写成「指标为行、战斗为列」的矩阵块。"""
from __future__ import annotations

from openpyxl.styles import Alignment, Font

from 战斗模拟.config import DATA_START, HEADER_ROW, SHEET_TASK
from 战斗模拟.load.blocks import block_spans
from 战斗模拟.sim.metrics import catalog_map

KEY_FIELDS = ("层级", "主体类型", "主体名", "指标名", "单位", "口径", "用途")
MATRIX_BLOCKS = {
    "对局": "战斗矩阵·对局",
    "阵营": "战斗矩阵·对局",
    "单位": "战斗矩阵·单位",
    "技能": "战斗矩阵·技能效果",
    "效果": "战斗矩阵·技能效果",
}


def _headers(ws, start: int, end: int) -> dict[str, int]:
    out = {}
    for col in range(start, end + 1):
        name = ws.cell(HEADER_ROW, col).value
        if name:
            out[str(name).strip()] = col
    return out


def _block(ws, name: str):
    for title, start, end in block_spans(ws):
        if title == name:
            return start, end
    return None


def _row_key(subject: str, subject_name: str, metric: str) -> tuple:
    return (str(subject or ""), str(subject_name or ""), str(metric or ""))


def _ensure_key_rows(ws, start: int, rows: list[dict]) -> dict[tuple, int]:
    """保证键列行存在，返回 (主体,主体名,指标名) -> 行号。"""
    headers = _headers(ws, start, start + len(KEY_FIELDS) - 1)
    name_col = headers.get("主体名", start + 2)
    metric_col = headers.get("指标名", start + 3)
    kind_col = headers.get("主体类型", start + 1)
    index = {}
    last = DATA_START - 1
    empty = 0
    for row in range(DATA_START, DATA_START + 5000):
        kind = ws.cell(row, kind_col).value
        subj = ws.cell(row, name_col).value
        metric = ws.cell(row, metric_col).value
        if kind in (None, "") and subj in (None, "") and metric in (None, ""):
            empty += 1
            if empty > 8:
                break
            continue
        empty = 0
        last = row
        index[_row_key(kind, subj, metric)] = row

    font = Font(name="微软雅黑", size=10)
    align = Alignment(vertical="center", wrap_text=True)
    meta = catalog_map()
    for item in rows:
        key = _row_key(item.get("主体"), item.get("主体名"), item.get("指标名"))
        if key in index:
            continue
        last += 1
        index[key] = last
        info = meta.get(str(item.get("指标名") or ""), None)
        values = {
            "层级": item.get("层级") or (info[1] if info else ""),
            "主体类型": item.get("主体"),
            "主体名": item.get("主体名"),
            "指标名": item.get("指标名"),
            "单位": item.get("单位") or (info[3] if info else ""),
            "口径": item.get("口径") or (info[4] if info else ""),
            "用途": item.get("用途") or (info[5] if info else ""),
        }
        for field, value in values.items():
            col = headers.get(field)
            if not col:
                continue
            cell = ws.cell(last, col)
            cell.value = value
            cell.font = font
            cell.alignment = align
    return index


def _next_fight_col(ws, start: int, end: int) -> int:
    headers = _headers(ws, start, end)
    key_cols = {headers[name] for name in KEY_FIELDS if name in headers}
    first_data = (max(key_cols) + 1) if key_cols else start + len(KEY_FIELDS)
    col = first_data
    while col <= end and ws.cell(HEADER_ROW, col).value not in (None, ""):
        col += 1
    return col


def _grow_block(ws, block_name: str, amount: int = 64) -> tuple[int, int] | None:
    """块满时在下一块标题前插入空列，扩大本块。"""
    spans = block_spans(ws)
    for i, (title, start, end) in enumerate(spans):
        if title != block_name:
            continue
        insert_at = end + 1
        ws.insert_cols(insert_at, amount)
        return start, end + amount
    return None


def append_matrix_column(wb, rows: list[dict], column_title: str) -> dict[str, int]:
    """按主体类型分到三块矩阵，各追加一列。返回块名→列号。"""
    ws = wb[SHEET_TASK]
    buckets = {"战斗矩阵·对局": [], "战斗矩阵·单位": [], "战斗矩阵·技能效果": []}
    for row in rows:
        block = MATRIX_BLOCKS.get(str(row.get("主体") or ""))
        if block:
            buckets[block].append(row)

    written = {}
    font = Font(name="微软雅黑", size=10)
    for block_name, items in buckets.items():
        span = _block(ws, block_name)
        if span is None:
            continue
        start, end = span
        if not items:
            continue
        index = _ensure_key_rows(ws, start, items)
        col = _next_fight_col(ws, start, end)
        while col > end:
            grown = _grow_block(ws, block_name)
            if grown is None:
                break
            start, end = grown
            col = _next_fight_col(ws, start, end)
        if col > end:
            continue
        cell = ws.cell(HEADER_ROW, col)
        cell.value = column_title
        cell.font = font
        for item in items:
            key = _row_key(item.get("主体"), item.get("主体名"), item.get("指标名"))
            row_i = index.get(key)
            if row_i is None:
                continue
            value = item.get("均值")
            ws.cell(row_i, col).value = value
            ws.cell(row_i, col).font = font
        written[block_name] = col
    return written


def write_matrices(path: str, rows: list[dict], column_title: str) -> dict[str, int]:
    import openpyxl
    wb = openpyxl.load_workbook(path)
    try:
        written = append_matrix_column(wb, rows, column_title)
        wb.save(path)
        return written
    finally:
        wb.close()
