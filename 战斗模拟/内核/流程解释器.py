"""战斗流程解释器：按管线名加载 sheet 步骤并执行中文 DSL（含光环/驱散真执行）。

- 管线名驱动：新管线写入「战斗流程」后自动可列/可跑，无需改代码注册。
- 跳转：成功跳转/失败跳转；0=顺序下一，>0=同管线步骤序，-1=结束。
- 禁止 调用结算。
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable

from 战斗模拟.load.战斗流程 import (
    NEW_PIPE_ALIASES,
    加载战斗流程,
    加载事件接入,
    加载机制旋钮,
    加载机制流程,
    加载阶段规划,
    流程步骤,
    流程阶段,
)
from 战斗模拟.领域 import 流程语法 as 语法
from 战斗模拟.管道 import 护盾吸收 as _盾

# ---------------------------------------------------------------------------
# 错误
# ---------------------------------------------------------------------------


class 解释器错误(RuntimeError):
    """流程解释或 DSL 求值失败。"""


# ---------------------------------------------------------------------------
# 词法 / 求值辅助
# ---------------------------------------------------------------------------

_NUM_RE = re.compile(r"\d+(?:\.\d+)?")
_IDENT_START = re.compile(r"[\u4e00-\u9fffA-Za-z_]")


def _compact(s: str) -> str:
    return re.sub(r"\s+", "", str(s).strip())


def _match_call_arg(s: str, start: int) -> tuple[str, int]:
    if start >= len(s) or s[start] != "(":
        raise 解释器错误(f"期望 '(' 于 {start}: {s[:80]}")
    depth = 0
    i = start
    while i < len(s):
        ch = s[i]
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
            if depth == 0:
                return s[start + 1 : i], i + 1
        i += 1
    raise 解释器错误(f"括号不配对：{s[:80]}")


def _split_top_args(arg: str) -> list[str]:
    parts: list[str] = []
    depth = 0
    last = 0
    for i, ch in enumerate(arg):
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        elif ch == "," and depth == 0:
            parts.append(arg[last:i].strip())
            last = i + 1
    parts.append(arg[last:].strip())
    return [p for p in parts if p is not None]


def _is_whole_call(s: str, fname: str) -> tuple[bool, str]:
    """若 s 整体恰为 fname(args)，返回 (True, args)。"""
    prefix = fname + "("
    if not s.startswith(prefix):
        return False, ""
    arg, end = _match_call_arg(s, len(fname))
    if end != len(s):
        return False, ""
    return True, arg


# 设* → 上下文键
_SET_FUNCS = frozenset(语法.设值函数)


def _set_key(fname: str) -> str:
    if not fname.startswith("设"):
        raise 解释器错误(f"非设值函数: {fname}")
    return fname[1:]


# ---------------------------------------------------------------------------
# 表达式求值器
# ---------------------------------------------------------------------------


class _ExprEval:
    """递归下降：算术 / 比较 / || && / 函数调用。"""

    def __init__(self, ctx: dict[str, Any], interp: "流程解释器") -> None:
        self.ctx = ctx
        self.interp = interp
        self.s = ""
        self.i = 0

    def eval(self, expr: str) -> Any:
        self.s = _compact(expr)
        self.i = 0
        if not self.s:
            raise 解释器错误("空表达式")
        v = self._parse_or()
        if self.i < len(self.s):
            raise 解释器错误(
                f"表达式未消费完 @{self.i}: {self.s[self.i:self.i+40]} ← {self.s[:80]}"
            )
        return v

    def _peek(self) -> str:
        return self.s[self.i] if self.i < len(self.s) else ""

    def _parse_or(self) -> Any:
        left = self._parse_and()
        while self.s.startswith("||", self.i):
            self.i += 2
            right = self._parse_and()
            left = bool(left) or bool(right)
        return left

    def _parse_and(self) -> Any:
        left = self._parse_compare()
        while self.s.startswith("&&", self.i):
            self.i += 2
            right = self._parse_compare()
            left = bool(left) and bool(right)
        return left

    def _parse_compare(self) -> Any:
        left = self._parse_add()
        for op in ("==", "!=", "<=", ">=", "<", ">"):
            if self.s.startswith(op, self.i):
                self.i += len(op)
                right = self._parse_add()
                if op == "==":
                    return self._eq(left, right)
                if op == "!=":
                    return not self._eq(left, right)
                lf, rf = float(left), float(right)
                if op == "<":
                    return lf < rf
                if op == ">":
                    return lf > rf
                if op == "<=":
                    return lf <= rf
                if op == ">=":
                    return lf >= rf
        return left

    @staticmethod
    def _eq(a: Any, b: Any) -> bool:
        if isinstance(a, str) or isinstance(b, str):
            return str(a) == str(b)
        try:
            return float(a) == float(b)
        except (TypeError, ValueError):
            return a == b

    def _parse_add(self) -> Any:
        left = self._parse_mul()
        while True:
            ch = self._peek()
            if ch == "+":
                self.i += 1
                right = self._parse_mul()
                left = float(left) + float(right)
            elif ch == "-":
                self.i += 1
                right = self._parse_mul()
                left = float(left) - float(right)
            else:
                break
        return left

    def _parse_mul(self) -> Any:
        left = self._parse_unary()
        while True:
            ch = self._peek()
            if ch == "*":
                self.i += 1
                right = self._parse_unary()
                left = float(left) * float(right)
            elif ch == "/":
                self.i += 1
                right = self._parse_unary()
                denom = float(right)
                left = float(left) / denom if denom != 0 else 0.0
            else:
                break
        return left

    def _parse_unary(self) -> Any:
        if self._peek() == "-":
            self.i += 1
            return -float(self._parse_unary())
        if self._peek() == "+":
            self.i += 1
            return float(self._parse_unary())
        return self._parse_primary()

    def _parse_primary(self) -> Any:
        ch = self._peek()
        if ch == "(":
            self.i += 1
            v = self._parse_or()
            if self._peek() != ")":
                raise 解释器错误(f"缺少 ')'：{self.s[self.i:self.i+40]}")
            self.i += 1
            return v
        if ch.isdigit() or (ch == "." and self.i + 1 < len(self.s) and self.s[self.i + 1].isdigit()):
            m = _NUM_RE.match(self.s, self.i)
            if not m:
                raise 解释器错误(f"非法数字：{self.s[self.i:self.i+8]}")
            self.i = m.end()
            return float(m.group(0))
        # 标识符 / 函数
        if not _IDENT_START.match(ch or ""):
            raise 解释器错误(f"意外字符 {ch!r} @{self.i}: {self.s[:80]}")
        start = self.i
        self.i += 1
        while self.i < len(self.s):
            c = self.s[self.i]
            if _IDENT_START.match(c) or c.isdigit() or c in "%_":
                self.i += 1
            else:
                break
        name = self.s[start : self.i]
        if self._peek() == "(":
            arg, end = _match_call_arg(self.s, self.i)
            self.i = end
            return self.interp._call_func(name, arg, self.ctx)
        return self.interp._resolve_ident(name, self.ctx)


# ---------------------------------------------------------------------------
# 上下文
# ---------------------------------------------------------------------------

_DEFAULT_LOCALS: dict[str, Any] = {
    "管线伤害": 0.0,
    "基础伤害": 0.0,
    "最终伤害": 0.0,
    "治疗值": 0.0,
    "治疗量": 0.0,
    "基础治疗": 0.0,
    "格挡减免量": 0.0,
    "反伤量": 0.0,
    "吸血量": 0.0,
    "仇恨系数": 1.0,
    "仇恨基础比例": 0.05,
    "模式": "PVE",
    "伤害类型": "物理",
    "伤害来源": "直接",
    "治疗来源": "直接",
    "当前生命": 0.0,
    "结算伤害": 0.0,
    "结算目标": "守方",
    "结算目标生命": 0.0,
    "护盾吸收量": 0.0,
    "护盾游标": 0.0,
    "护盾层数": 0.0,
    "当前护盾": None,
    "本层吸收": 0.0,
    "过量伤害": 0.0,
    "本次破盾数": 0.0,
    "生命空档": 0.0,
    "实际治疗": 0.0,
    "溢出治疗": 0.0,
    "过量治疗": 0.0,
    "过量转盾量": 0.0,
    "治疗吸收量": 0.0,
    "治疗吸收游标": 0.0,
    "治疗吸收层数": 0.0,
    "当前治疗吸收": None,
    "本层治疗吸收": 0.0,
    "破盾结束Buff列表": [],
    "对应吸血": 0.0,
    "对应反伤": 0.0,
    "实际扣血": 0.0,
    "实际反伤扣血": 0.0,
    "伤害仇恨": 0.0,
    "治疗仇恨": 0.0,
    "命中位": "命中判定",
    "暴击位": "可暴击",
    "异常挂载结果": 0.0,
    "异常层数": 0.0,
    "剩余时长": 0.0,
    "基础时长": 0.0,
    "基础层数": 1.0,
    "最大层数": 1.0,
    "当前层数": 0.0,
    "附着概率": 1.0,
    "效果代号": "",
    "效果类型": "减益",
    "叠加规则": "独立",
    "刷新规则": "重置刷新",
    "时长继承": 0.0,
    "跳动间隔": 0.0,
    "生效时机": "立即",
    "互斥组": "空",
    "控制递减组": "空",
    "递减系数": 1.0,
    "控制行为": "",
    "首领抗性": 0.0,
    "数值类型": "无",
    "异常跳伤基础": 0.0,
    "异常跳疗基础": 0.0,
    "跳伤伤害": 0.0,
    "跳疗治疗": 0.0,
    "旧强度": 0.0,
    "新强度": 0.0,
    "PVP时长系数": 1.0,
    "叠加动作": "",
    "标签": "",
    # Phase A 异常细化 / 驱散
    "快照时机": "无",
    "快照属性列表": "",
    "层衰减间隔": 0.0,
    "层衰减数量": 0.0,
    "驱散强度": 0.0,
    "已驱散数": 0.0,
    "驱散结果": 0.0,
    "驱散优先级": "",
    "驱散类型": "",
    "驱散类型等级": "可驱散",
    "可窃取": "否",
    "到期动作": "",
    "吟唱时长": 0.0,
    "触发结果": 0.0,
    "触发概率": 0.0,
    "通知结果": 0.0,
}


def _panel_get(panel: dict[str, Any] | None, key: str, default: float = 0.0) -> float:
    if not panel:
        return float(default)
    nested = panel.get("面板") if isinstance(panel.get("面板"), dict) else None
    for src in (panel, nested):
        if not src:
            continue
        if key in src and src[key] is not None:
            try:
                return float(src[key])
            except (TypeError, ValueError):
                pass
    return float(default)


def _side_states(side: dict[str, Any] | None) -> set[str]:
    if not side:
        return set()
    raw = side.get("状态") or side.get("states") or set()
    if isinstance(raw, set):
        return set(raw)
    if isinstance(raw, (list, tuple)):
        return {str(x) for x in raw}
    return set()


def _side_flags(side: dict[str, Any] | None) -> set[str]:
    if not side:
        return set()
    raw = side.get("标记") or side.get("flags") or set()
    if isinstance(raw, set):
        return set(raw)
    if isinstance(raw, (list, tuple)):
        return {str(x) for x in raw}
    return set()


def 构建上下文(
    *,
    基础伤害: float = 0.0,
    基础治疗: float = 0.0,
    伤害类型: str = "物理",
    伤害来源: str = "直接",
    治疗来源: str = "直接",
    攻方: dict[str, Any] | None = None,
    守方: dict[str, Any] | None = None,
    rng: Any = None,
    仇恨系数: float = 1.0,
    标记: Iterable[str] | None = None,
    护盾层: list[dict[str, Any]] | None = None,
    效果覆盖: dict[str, float] | None = None,
    **extra: Any,
) -> dict[str, Any]:
    """合成管线上下文（mock 面板友好）。"""
    # 使用调用方面板引用（就地写回生命/护盾）；未传则新建空字典
    atk = 攻方 if isinstance(攻方, dict) else {}
    dfd = 守方 if isinstance(守方, dict) else {}
    # 守方护盾默认；若调用方显式传 护盾层 则用之，并写回守方面板以保持引用一致
    if 护盾层 is not None:
        shields = list(护盾层)
        dfd["护盾层"] = shields
    else:
        if not isinstance(dfd.get("护盾层"), list):
            dfd["护盾层"] = list(dfd.get("护盾层") or [])
        shields = dfd["护盾层"]
    if not isinstance(atk.get("护盾层"), list):
        atk["护盾层"] = list(atk.get("护盾层") or [])
    if not isinstance(dfd.get("治疗吸收层"), list):
        dfd["治疗吸收层"] = list(dfd.get("治疗吸收层") or [])
    if not isinstance(atk.get("治疗吸收层"), list):
        atk["治疗吸收层"] = list(atk.get("治疗吸收层") or [])
    heal_abs = dfd["治疗吸收层"]
    ctx: dict[str, Any] = dict(_DEFAULT_LOCALS)
    ctx.update(
        {
            "基础伤害": float(基础伤害),
            "管线伤害": float(基础伤害),
            "基础治疗": float(基础治疗),
            "治疗值": float(基础治疗),
            "伤害类型": str(伤害类型),
            "伤害来源": str(伤害来源),
            "治疗来源": str(治疗来源),
            "仇恨系数": float(仇恨系数),
            "攻方": atk,
            "守方": dfd,
            "标记": set(标记 or []),
            "护盾层": shields,
            "护盾层数": float(len(shields)),
            "治疗吸收层": heal_abs,
            "治疗吸收层数": float(len(heal_abs)),
            "结算目标": "守方",
            "结算目标生命": float(dfd.get("生命") or dfd.get("当前生命") or 0.0),
            "当前生命": float(dfd.get("当前生命") or dfd.get("生命") or 0.0),
            "本次破盾数": 0.0,
            "rng": rng,
            "效果覆盖": dict(效果覆盖 or {}),
            "步骤轨迹": [],
        }
    )
    ctx.update(extra)
    return ctx


# ---------------------------------------------------------------------------
# 解释器
# ---------------------------------------------------------------------------

OpcodeHandler = Callable[[流程步骤, dict[str, Any], "流程解释器"], bool | None]


@dataclass
class 管线执行结果:
    管线: str
    最终伤害: float = 0.0
    实际扣血: float = 0.0
    实际治疗: float = 0.0
    伤害仇恨: float = 0.0
    治疗仇恨: float = 0.0
    异常挂载结果: float = 0.0
    标记: set[str] = field(default_factory=set)
    上下文: dict[str, Any] = field(default_factory=dict)
    步骤轨迹: list[dict[str, Any]] = field(default_factory=list)
    步数: int = 0


class 流程解释器:
    """加载战斗流程 sheet，按管线名执行。"""

    def __init__(
        self,
        管线表: dict[str, list[流程步骤]] | None = None,
        *,
        工作簿路径: str | Path | None = None,
        属性名集合: set[str] | None = None,
        校验公式: bool = False,
        最大步数: int = 10_000,
    ) -> None:
        self.最大步数 = 最大步数
        self.属性名集合 = 属性名集合
        self._opcode_handlers: dict[str, OpcodeHandler] = {}
        self._unknown_formulas: list[str] = []
        self.事件表 = {}
        self.阶段表: dict[str, list] = {}
        self.机制表: dict[str, list] = {}
        self.机制旋钮: dict[str, float] = {}
        if 管线表 is not None:
            self.管线表 = 管线表
        elif 工作簿路径 is not None:
            self._装载工作簿(工作簿路径, 校验公式=校验公式)
        else:
            self.管线表 = {}

    # -- 加载 ----------------------------------------------------------------

    @classmethod
    def 从工作簿加载(
        cls,
        path: str | Path,
        *,
        校验公式: bool = False,
        属性名集合: set[str] | None = None,
    ) -> dict[str, list[流程步骤]]:
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True)
        try:
            pipes = 加载战斗流程(wb)
            attrs = 属性名集合
            if 校验公式:
                if attrs is None:
                    try:
                        attrs = 语法.从属性总表加载用途名(wb)
                    except Exception:
                        attrs = None
                for name, steps in pipes.items():
                    for st in steps:
                        try:
                            语法.校验判定公式(st.判定公式, attrs)
                        except 语法.流程语法错误 as e:
                            raise 解释器错误(
                                f"管线「{name}」步骤{st.步骤序} 公式非法: {e}"
                            ) from e
            return pipes
        finally:
            wb.close()

    def _装载工作簿(self, path: str | Path, *, 校验公式: bool = False) -> None:
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True)
        try:
            self.管线表 = 加载战斗流程(wb)
            self.阶段表 = 加载阶段规划(wb)
            self.机制表 = 加载机制流程(wb)
            self.机制旋钮 = 加载机制旋钮(wb)
            self.事件表 = 加载事件接入(wb)
            if 校验公式:
                attrs = self.属性名集合
                if attrs is None:
                    try:
                        attrs = 语法.从属性总表加载用途名(wb)
                    except Exception:
                        attrs = None
                for name, steps in self.管线表.items():
                    for st in steps:
                        try:
                            语法.校验判定公式(st.判定公式, attrs)
                        except 语法.流程语法错误 as e:
                            raise 解释器错误(
                                f"管线「{name}」步骤{st.步骤序} 公式非法: {e}"
                            ) from e
        finally:
            wb.close()

    def _起始下标(self, 管线名: str, steps, context: dict) -> int:
        from 战斗模拟.内核.事件 import 规范化事件种类

        by_order = {s.步骤序: s for s in steps}
        order_list = [s.步骤序 for s in steps]
        entry = 0
        raw = context.get("入口步骤")
        try:
            if raw not in (None, ""):
                entry = int(raw)
        except (TypeError, ValueError):
            entry = 0
        row = None
        kind = 规范化事件种类(context.get("事件种类"))
        if kind:
            row = (self.事件表 or {}).get(kind)
        if entry <= 0 and row is not None and row.启用 and row.入口步骤 > 0:
            entry = int(row.入口步骤)
        if row is not None and row.启用 and row.来源覆盖 and not context.get("_来源已覆盖"):
            src = row.来源覆盖
            if "伤害" in 管线名:
                context["伤害来源"] = src
            if "治疗" in 管线名:
                context["治疗来源"] = src
            context["_来源已覆盖"] = True
        if entry > 0:
            if entry not in by_order:
                raise 解释器错误(
                    f"管线「{管线名}」入口步骤 {entry} 不存在（事件种类={kind or '—'}）"
                )
            return order_list.index(entry)
        return 0

    def 管线名列表(self) -> list[str]:
        names = list(self.管线表.keys())
        for n in list(self.阶段表) + list(self.机制表):
            if n not in names:
                names.append(n)
        return names

    def _映射规划路名(self, 管线名: str, ctx: dict[str, Any] | None = None) -> str:
        if 管线名 in self.阶段表 or 管线名 in self.机制表:
            return 管线名
        from 战斗模拟.load.战斗流程 import MODE_PIPE

        mode = self._当前模式(ctx or {}, 管线名)
        keyed = MODE_PIPE.get((管线名, mode))
        legacy = {
            "伤害PVE": "伤害环境",
            "伤害PVP": "伤害对战",
            "治疗PVE": "治疗环境",
            "治疗PVP": "治疗对战",
            "状态PVE": "状态环境",
            "状态PVP": "状态对战",
        }
        for cand in (
            管线名,
            keyed,
            NEW_PIPE_ALIASES.get(管线名),
            NEW_PIPE_ALIASES.get(keyed or ""),
            legacy.get(管线名),
            legacy.get(keyed or ""),
        ):
            if cand and (cand in self.阶段表 or cand in self.机制表):
                return cand
        return 管线名

    def _当前模式(self, ctx: dict[str, Any], 管线名: str = "") -> str:
        raw = str(ctx.get("模式") or "").strip().upper()
        if raw in {"PVP", "战场"}:
            return "PVP"
        if raw in {"PVE", "副本"}:
            return "PVE"
        name = str(管线名)
        if "对战" in name or "PVP" in name:
            return "PVP"
        if "环境" in name or "PVE" in name:
            return "PVE"
        return "PVE"

    def _适用本模式(self, 适用: str, mode: str) -> bool:
        text = str(适用 or "共用").strip()
        if text in {"", "都用", "全部", "共用"}:
            return True
        if text in {"仅副本", "副本"}:
            return mode == "PVE"
        if text in {"仅战场", "战场"}:
            return mode == "PVP"
        return True

    def 注册操作码(self, 步骤名: str, handler: OpcodeHandler) -> None:
        self._opcode_handlers[步骤名] = handler

    def _挂钩特殊事件(self, 段名: str, context: dict[str, Any]) -> None:
        if "特殊事件" not in self.阶段表:
            return
        entry = {
            "写出最终伤害": 1,
            "应用受治疗": 4,
            "执行扣生命": 6,
        }.get(段名)
        if not entry:
            return
        if entry == 1 and not (
            self._求值语句("有状态(无敌)||结算目标有状态(无敌)", context, 段名="特殊事件")
        ):
            return
        if entry == 4 and not (
            self._求值语句("有状态(禁止受疗)||有状态(无法被治疗)", context, 段名="特殊事件")
        ):
            return
        if entry == 6 and not (
            self._求值语句(
                "有状态(锁1血)||有标记(锁1血)||有状态(不死)||有标记(不死)",
                context,
                段名="特殊事件",
            )
        ):
            return
        prev = context.get("入口步骤")
        context["入口步骤"] = entry
        try:
            self._执行阶段管线("特殊事件", context, 请求名="特殊事件")
        finally:
            if prev in (None, ""):
                context.pop("入口步骤", None)
            else:
                context["入口步骤"] = prev

    # -- 执行 ----------------------------------------------------------------

    def 执行管线(
        self,
        管线名: str,
        ctx: dict[str, Any] | None = None,
        **kwargs: Any,
    ) -> 管线执行结果:
        from 战斗模拟.领域.流程语法 import 规范化管线名

        管线名 = 规范化管线名(管线名)
        context = ctx if ctx is not None else None
        planned = self._映射规划路名(管线名, context or {})
        if planned in self.阶段表 and self.阶段表[planned]:
            context = context if context is not None else 构建上下文(**kwargs)
            return self._执行阶段管线(planned, context, 请求名=管线名)
        if planned in self.机制表 and self.机制表[planned]:
            context = ctx if ctx is not None else 构建上下文(**kwargs)
            return self._执行阶段管线(planned, context, 请求名=管线名)
        if 管线名 not in self.管线表:
            known = ", ".join(self.管线名列表()) or "(无)"
            raise 解释器错误(f"管线「{管线名}」不存在；已加载: {known}")
        steps = self.管线表[管线名]
        if not steps:
            raise 解释器错误(f"管线「{管线名}」无步骤")
        context = ctx if ctx is not None else 构建上下文(**kwargs)
        if not context.get("结算目标"):
            context["结算目标"] = "守方"
        self._sync_settle_life_from_side(context)
        by_order = {s.步骤序: s for s in steps}
        order_list = [s.步骤序 for s in steps]
        idx = self._起始下标(管线名, steps, context)
        n_exec = 0
        trail: list[dict[str, Any]] = context.setdefault("步骤轨迹", [])

        while 0 <= idx < len(order_list):
            if n_exec >= self.最大步数:
                raise 解释器错误(f"管线「{管线名}」超过最大步数 {self.最大步数}")
            order = order_list[idx]
            step = by_order[order]
            n_exec += 1

            # 可选 opcode 覆写
            handler = self._opcode_handlers.get(step.步骤名)
            if handler is not None:
                ok = handler(step, context, self)
                if ok is None:
                    ok = True
            else:
                ok = self._执行步骤公式(step, context)

            trail.append(
                {
                    "步骤序": step.步骤序,
                    "步骤名": step.步骤名,
                    "成功": bool(ok),
                    "公式": step.判定公式,
                }
            )
            jump = step.成功跳转 if ok else step.失败跳转
            if jump == -1:
                break
            if jump == 0:
                idx += 1
                continue
            # 绝对步骤序
            if jump not in by_order:
                raise 解释器错误(
                    f"管线「{管线名}」步骤{step.步骤序} 跳转目标 {jump} 不存在"
                )
            idx = order_list.index(jump)

        # 回写实体生命
        self._写回生命(context)

        return 管线执行结果(
            管线=管线名,
            最终伤害=float(context.get("最终伤害") or 0.0),
            实际扣血=float(context.get("实际扣血") or 0.0),
            实际治疗=float(context.get("实际治疗") or 0.0),
            伤害仇恨=float(context.get("伤害仇恨") or 0.0),
            治疗仇恨=float(context.get("治疗仇恨") or 0.0),
            异常挂载结果=float(context.get("异常挂载结果") or 0.0),
            标记=set(context.get("标记") or []),
            上下文=context,
            步骤轨迹=list(trail),
            步数=n_exec,
        )

    def _执行阶段管线(
        self,
        路名: str,
        context: dict[str, Any],
        *,
        请求名: str = "",
        入口段名: str = "",
    ) -> 管线执行结果:
        if not context.get("结算目标"):
            context["结算目标"] = "守方"
        mode = self._当前模式(context, 请求名 or 路名)
        context["模式"] = mode
        for k, v in (self.机制旋钮 or {}).items():
            context.setdefault(k, v)
        self._sync_settle_life_from_side(context)
        raw_stages = self.阶段表.get(路名) or self.机制表.get(路名) or []
        stages = [s for s in raw_stages if s.启用 and self._适用本模式(s.适用, mode)]
        if not stages:
            raise 解释器错误(f"规划路「{路名}」无可用阶段")
        start_name = 入口段名 or str(context.get("入口段名") or "")
        entry_order = 0
        raw_entry = context.get("入口步骤")
        try:
            if raw_entry not in (None, ""):
                entry_order = int(raw_entry)
        except (TypeError, ValueError):
            entry_order = 0
        row = None
        from 战斗模拟.内核.事件 import 规范化事件种类
        kind = 规范化事件种类(context.get("事件种类"))
        if kind:
            row = (self.事件表 or {}).get(kind)
        if row is not None and row.启用:
            if not start_name:
                start_name = str(getattr(row, "入口段名", "") or "")
            if entry_order <= 0:
                entry_order = int(getattr(row, "入口步骤", 0) or 0)
            if row.来源覆盖 and not context.get("_来源已覆盖"):
                if 路名 in {"伤害", "伤害PVE", "伤害PVP", "伤害环境", "伤害对战"}:
                    context["伤害来源"] = row.来源覆盖
                if 路名 in {"治疗", "治疗PVE", "治疗PVP", "治疗环境", "治疗对战"}:
                    context["治疗来源"] = row.来源覆盖
                context["_来源已覆盖"] = True
        by_order = {st.第几段: i for i, st in enumerate(stages)}
        idx = 0
        if entry_order > 0 and entry_order in by_order:
            idx = by_order[entry_order]
        elif start_name:
            for i, st in enumerate(stages):
                if st.段名 == start_name:
                    idx = i
                    break
        n_exec = 0
        trail: list[dict[str, Any]] = context.setdefault("步骤轨迹", [])
        while 0 <= idx < len(stages):
            if n_exec >= self.最大步数:
                raise 解释器错误(f"规划路「{路名}」超过最大步数 {self.最大步数}")
            st = stages[idx]
            n_exec += 1
            skipped = False
            if st.跳过条件:
                try:
                    skipped = bool(self._求值语句(st.跳过条件, context))
                except 解释器错误:
                    skipped = False
            ok = True
            if not skipped:
                if st.判定公式:
                    ok = self._执行多句(st.判定公式, context, 段名=st.段名)
                for mech in [x.strip() for x in str(st.调用机制 or "").replace("、", "；").split("；") if x.strip()]:
                    self._执行机制(mech, context, mode=mode)
            trail.append(
                {
                    "步骤序": st.第几段,
                    "步骤名": st.段名,
                    "成功": bool(ok),
                    "公式": st.判定公式,
                    "跳过": skipped,
                    "机制": st.调用机制,
                    "跳转": (st.成功跳转 if ok else st.失败跳转) if not skipped else 0,
                }
            )
            if not skipped and 路名 != "特殊事件":
                self._挂钩特殊事件(st.段名, context)
            if skipped:
                idx += 1
                continue
            jump = st.成功跳转 if ok else st.失败跳转
            if jump == -1:
                break
            if jump == 0:
                idx += 1
                continue
            if jump not in by_order:
                raise 解释器错误(
                    f"规划路「{路名}」步骤{st.第几段} 跳转目标 {jump} 不存在"
                )
            idx = by_order[jump]
        self._写回生命(context)
        return 管线执行结果(
            管线=路名,
            最终伤害=float(context.get("最终伤害") or 0.0),
            实际扣血=float(context.get("实际扣血") or 0.0),
            实际治疗=float(context.get("实际治疗") or 0.0),
            伤害仇恨=float(context.get("伤害仇恨") or 0.0),
            治疗仇恨=float(context.get("治疗仇恨") or 0.0),
            异常挂载结果=float(context.get("异常挂载结果") or 0.0),
            标记=set(context.get("标记") or []),
            上下文=context,
            步骤轨迹=list(trail),
            步数=n_exec,
        )

    def _split_seq(self, s: str) -> list[str]:
        parts: list[str] = []
        depth = 0
        last = 0
        text = str(s or "")
        i = 0
        while i < len(text):
            ch = text[i]
            if ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
            elif depth == 0 and text[i] in "；;":
                part = text[last:i].strip()
                if part:
                    parts.append(part)
                last = i + 1
            i += 1
        tail = text[last:].strip()
        if tail:
            parts.append(tail)
        return parts

    def _执行多句(self, raw: str, ctx: dict[str, Any], *, 段名: str = "") -> bool:
        text = str(raw or "").strip()
        if not text:
            return True
        if text.startswith("若") and "则" in text:
            cond, then = text[1:].split("则", 1)
            if self._求值语句(cond.strip(), ctx, 段名=段名):
                return self._执行多句(then.strip(), ctx, 段名=段名)
            return False
        last = True
        for part in self._split_seq(text):
            last = bool(self._求值语句(part, ctx, 段名=段名))
        return last

    def _求值语句(self, raw: str, ctx: dict[str, Any], *, 段名: str = "") -> Any:
        s = str(raw or "").strip()
        if s.startswith("若") and "则" in s:
            body = s[1:]
            cond, then = body.split("则", 1)
            cond = cond.strip()
            then = then.strip()
            if self._求值语句(cond, ctx, 段名=段名):
                return self._执行多句(then, ctx, 段名=段名)
            return False
        step = 流程步骤(管线=段名 or "阶段", 步骤序=0, 步骤名=段名 or "阶段", 判定公式=s, 成功跳转=0, 失败跳转=0)
        try:
            return self._执行步骤公式(step, ctx)
        except 解释器错误 as e:
            # 策划说明句（非 DSL）不当作失败
            if any(x in s for x in ("从前往后", "读效果总表", "按优先级", "独立/", "层衰减间隔")):
                return True
            raise e

    def _执行机制(self, name: str, ctx: dict[str, Any], *, mode: str) -> None:
        key = str(name or "").strip()
        if not key:
            return
        handlers = {
            "护盾吸收": self._mech_护盾吸收,
            "治疗吸收": self._mech_治疗吸收,
            "过量转盾": self._mech_过量转盾,
            "锁1血": self._mech_锁1血,
            "吸血与反伤": self._mech_吸血与反伤,
            "仇恨": self._mech_仇恨,
            "光环挂载": self._mech_光环挂载,
            "光环存续": self._mech_光环存续,
            "驱散与偷取": self._mech_驱散,
            "控制递减": self._mech_控制递减,
            "控制命中": self._mech_控制命中,
            "立即跳动": self._mech_立即跳动,
        }
        fn = handlers.get(key)
        if fn is None:
            raise 解释器错误(f"未知机制「{key}」")
        fn(ctx)

    def _dsl(self, ctx: dict[str, Any], formula: str, name: str = "机制") -> bool:
        return bool(self._执行多句(formula, ctx, 段名=name))

    def _mech_护盾吸收(self, ctx: dict[str, Any]) -> None:
        self._dsl(ctx, "设护盾吸收量(0)；设护盾游标(0)", "护盾吸收")
        ctx["护盾层数"] = float(len(ctx.get("护盾层") or []))
        while float(ctx.get("护盾游标") or 0) < float(ctx.get("护盾层数") or 0):
            self._dsl(ctx, "设当前护盾(取护盾层(护盾游标))", "护盾吸收")
            ok = self._求值语句("护盾类型可吸收(当前护盾,伤害类型)", ctx, 段名="护盾吸收")
            if ok:
                self._dsl(
                    ctx,
                    "设本层吸收(最小(结算伤害,护盾剩余容量(当前护盾)))；扣减护盾层(当前护盾,本层吸收)；设护盾吸收量(护盾吸收量+本层吸收)；设结算伤害(最大(0,结算伤害-本层吸收))",
                    "护盾吸收",
                )
            self._dsl(ctx, "设护盾游标(护盾游标+1)", "护盾吸收")
        if float(ctx.get("护盾吸收量") or 0) > 0:
            self._dsl(ctx, "标记(已破盾)", "护盾吸收")
        self._dsl(ctx, "设过量伤害(最大(0,结算伤害))", "护盾吸收")

    def _mech_锁1血(self, ctx: dict[str, Any]) -> None:
        self._dsl(ctx, "设当前生命(最大(0,当前生命-结算伤害))", "锁1血")
        if self._求值语句("有状态(锁1血)||有标记(锁1血)", ctx, 段名="锁1血"):
            self._dsl(ctx, "设当前生命(最大(1,当前生命))", "锁1血")
        self._dsl(ctx, "设实际扣血(最大(0,结算伤害))", "锁1血")
        if self._求值语句("当前生命<=0", ctx, 段名="锁1血"):
            self._dsl(ctx, "标记(已死亡)", "锁1血")
        self._dsl(ctx, "设结算目标生命(当前生命)", "锁1血")

    def _mech_治疗吸收(self, ctx: dict[str, Any]) -> None:
        if float(ctx.get("治疗值") or 0) <= 0:
            self._dsl(ctx, "设实际治疗(0)", "治疗吸收")
            return
        self._dsl(ctx, "设治疗吸收量(0)；设治疗吸收游标(0)", "治疗吸收")
        ctx["治疗吸收层数"] = float(len(ctx.get("治疗吸收层") or []))
        while float(ctx.get("治疗吸收游标") or 0) < float(ctx.get("治疗吸收层数") or 0):
            self._dsl(ctx, "设当前治疗吸收(取治疗吸收层(治疗吸收游标))", "治疗吸收")
            if self._求值语句("治疗吸收剩余容量(当前治疗吸收)>0", ctx, 段名="治疗吸收"):
                self._dsl(
                    ctx,
                    "设本层治疗吸收(最小(治疗值,治疗吸收剩余容量(当前治疗吸收)))；扣减治疗吸收层(当前治疗吸收,本层治疗吸收)；设治疗吸收量(治疗吸收量+本层治疗吸收)",
                    "治疗吸收",
                )
            self._dsl(ctx, "设治疗吸收游标(治疗吸收游标+1)", "治疗吸收")
        self._dsl(
            ctx,
            "设治疗值(最大(0,治疗值-治疗吸收量))；设生命空档(最大(0,结算目标属性(生命值)-当前生命))；设实际治疗(最小(治疗值,生命空档))；设过量治疗(最大(0,治疗值-实际治疗))；设当前生命(当前生命+实际治疗)；设结算目标生命(当前生命)",
            "治疗吸收",
        )
        self._mech_过量转盾(ctx)

    def _mech_过量转盾(self, ctx: dict[str, Any]) -> None:
        if float(ctx.get("过量治疗") or 0) <= 0:
            return
        if "过量转盾开启" not in (ctx.get("标记") or set()):
            return
        self._dsl(ctx, "执行过量转盾(过量治疗)；设过量转盾量(过量治疗)", "过量转盾")

    def _mech_吸血与反伤(self, ctx: dict[str, Any]) -> None:
        self._dsl(
            ctx,
            "设吸血量(最终伤害*((伤害类型==物理)*效果(物理吸血效果,攻方)+(伤害类型==魔法)*效果(魔法吸血效果,攻方)))；设反伤量(最终伤害*((伤害类型==物理)*效果(物理反伤效果,守方)+(伤害类型==魔法)*效果(魔法反伤效果,守方)))",
            "吸血与反伤",
        )
        if float(ctx.get("吸血量") or 0) > 0:
            self._dsl(ctx, "设结算目标(攻方)；设当前生命(当前生命+吸血量)；设结算目标(守方)", "吸血与反伤")
        if float(ctx.get("反伤量") or 0) > 0:
            self._dsl(
                ctx,
                "标记(反伤扣血中)；设结算目标(攻方)；设当前生命(最大(0,当前生命-反伤量))；设实际反伤扣血(反伤量)；取消标记(反伤扣血中)",
                "吸血与反伤",
            )
            if self._求值语句("当前生命<=0", ctx, 段名="吸血与反伤"):
                self._dsl(ctx, "标记(攻方已死亡)", "吸血与反伤")
            self._dsl(ctx, "设结算目标(守方)", "吸血与反伤")

    def _mech_仇恨(self, ctx: dict[str, Any]) -> None:
        if float(ctx.get("实际扣血") or 0) > 0 or float(ctx.get("反伤量") or 0) > 0:
            if "已死亡" not in (ctx.get("标记") or set()):
                self._dsl(ctx, "设伤害仇恨((实际扣血+反伤量)*仇恨系数)", "仇恨")
            else:
                self._dsl(ctx, "设伤害仇恨(0)", "仇恨")
        if float(ctx.get("实际治疗") or 0) > 0:
            self._dsl(ctx, "设治疗仇恨(实际治疗*仇恨基础比例*仇恨系数)", "仇恨")
        else:
            if ctx.get("治疗值") is not None and float(ctx.get("实际治疗") or 0) <= 0:
                ctx.setdefault("治疗仇恨", 0.0)

    def _mech_光环挂载(self, ctx: dict[str, Any]) -> None:
        if "施加已拒绝" in (ctx.get("标记") or set()):
            self._dsl(ctx, "设异常挂载结果(0)", "光环挂载")
            return
        code = ctx.get("效果代号") or ""
        if not self._求值语句("有同效果(效果代号)", ctx, 段名="光环挂载"):
            self._dsl(
                ctx,
                "设异常层数(基础层数)；设剩余时长(基础时长)；挂载效果(效果代号,异常层数,剩余时长)",
                "光环挂载",
            )
            return
        rule = str(ctx.get("叠加规则") or "叠层")
        if rule == "独立":
            self._dsl(ctx, "设异常层数(基础层数)；设剩余时长(基础时长)；挂载效果(效果代号,异常层数,剩余时长)", "光环挂载")
        elif rule == "替换":
            self._dsl(ctx, "移除同效果(效果代号)；设异常层数(基础层数)；设剩余时长(基础时长)；挂载效果(效果代号,异常层数,剩余时长)", "光环挂载")
        elif rule == "叠层":
            self._dsl(ctx, "设异常层数(最小(最大层数,当前层数+基础层数))", "光环挂载")
            refresh = str(ctx.get("刷新规则") or "重置刷新")
            if refresh == "继承刷新":
                self._dsl(ctx, "设剩余时长(最大(剩余时长,基础时长))", "光环挂载")
            else:
                self._dsl(ctx, "设剩余时长(基础时长)", "光环挂载")
            self._dsl(ctx, "挂载效果(效果代号,异常层数,剩余时长)", "光环挂载")
        else:
            self._dsl(ctx, "设异常层数(当前层数)；设剩余时长(剩余时长)；挂载效果(效果代号,异常层数,剩余时长)", "光环挂载")
        _ = code

    def _mech_光环存续(self, ctx: dict[str, Any]) -> None:
        if "层衰减结算" in (ctx.get("标记") or set()):
            if float(ctx.get("层衰减间隔") or 0) > 0:
                self._dsl(ctx, "设异常层数(最大(0,当前层数-层衰减数量))", "光环存续")
                if float(ctx.get("异常层数") or 0) <= 0:
                    self._dsl(ctx, "移除效果(效果代号)", "光环存续")
            return
        if "异常到期结算" in (ctx.get("标记") or set()):
            if str(ctx.get("到期动作") or "空") not in {"", "空"}:
                self._dsl(ctx, "执行到期动作(到期动作)", "光环存续")
            self._dsl(ctx, "移除效果(效果代号)", "光环存续")

    def _mech_驱散(self, ctx: dict[str, Any]) -> None:
        self._dsl(ctx, "设驱散结果(0)；设已驱散数(0)；设驱散强度(最大(0,驱散强度))", "驱散")
        guard = 0
        while float(ctx.get("已驱散数") or 0) < float(ctx.get("驱散强度") or 0) and guard < 64:
            guard += 1
            if not self._求值语句("有可驱散效果()", ctx, 段名="驱散"):
                break
            self._dsl(ctx, "选取可驱散效果(驱散优先级)", "驱散")
            if not self._求值语句("匹配驱散类型(标签,驱散类型)", ctx, 段名="驱散"):
                continue
            if self._求值语句("有标记(首领)&&驱散类型等级==不可驱散", ctx, 段名="驱散"):
                continue
            if str(ctx.get("驱散类型等级") or "") == "仅强驱" and "强驱" not in (ctx.get("标记") or set()):
                continue
            if self._求值语句("可窃取==是||有标记(可窃取技能)", ctx, 段名="驱散"):
                self._dsl(ctx, "转移效果(效果代号,攻方)", "驱散")
            else:
                self._dsl(ctx, "移除效果(效果代号)", "驱散")
            if str(ctx.get("到期动作") or "空") not in {"", "空"}:
                self._dsl(ctx, "执行到期动作(到期动作)", "驱散")
            self._dsl(ctx, "设已驱散数(已驱散数+1)；设驱散结果(已驱散数)", "驱散")
        self._dsl(ctx, "设驱散结果(已驱散数)", "驱散")

    def _mech_控制递减(self, ctx: dict[str, Any]) -> None:
        if str(ctx.get("效果类型") or "") != "控制":
            return
        self._dsl(ctx, "设剩余时长(剩余时长*结算目标递减系数(控制递减组))", "控制递减")
        if self._当前模式(ctx) == "PVP":
            self._dsl(ctx, "设剩余时长(剩余时长*PVP时长系数)", "控制递减")

    def _mech_控制命中(self, ctx: dict[str, Any]) -> None:
        if "施加已拒绝" in (ctx.get("标记") or set()):
            return
        if self._求值语句("随机()>=附着概率", ctx, 段名="控制命中"):
            self._dsl(ctx, "标记(施加已拒绝)；设异常挂载结果(0)", "控制命中")
            return
        if str(ctx.get("效果类型") or "") == "控制":
            if self._求值语句("随机()>=最大(0,最小(效果(控制效果),1))", ctx, 段名="控制命中"):
                self._dsl(ctx, "标记(施加已拒绝)；设异常挂载结果(0)", "控制命中")

    def _mech_立即跳动(self, ctx: dict[str, Any]) -> None:
        if "施加已拒绝" in (ctx.get("标记") or set()):
            return
        if str(ctx.get("生效时机") or "") == "立即":
            self._dsl(ctx, "标记(需立即跳转)", "立即跳动")
        if float(ctx.get("异常跳伤基础") or 0) > 0:
            self._dsl(ctx, "设跳伤伤害(异常跳伤基础)；设伤害来源(持续)", "立即跳动")
        if float(ctx.get("异常跳疗基础") or 0) > 0:
            self._dsl(ctx, "设跳疗治疗(异常跳疗基础)；设治疗来源(持续)", "立即跳动")

    def _执行步骤公式(self, step: 流程步骤, ctx: dict[str, Any]) -> bool:
        raw = step.判定公式
        if raw is None or not str(raw).strip():
            raise 解释器错误(f"步骤{step.步骤序}「{step.步骤名}」判定公式为空")
        s = _compact(raw)

        if "调用结算" in s:
            raise 解释器错误(
                f"步骤{step.步骤序} 使用已禁用的 调用结算：{s[:80]}"
            )

        # 设*
        for fname in _SET_FUNCS:
            ok, arg = _is_whole_call(s, fname)
            if not ok:
                continue
            return self._do_set(fname, arg, ctx)

        # 标记 / 取消标记
        ok, arg = _is_whole_call(s, "标记")
        if ok:
            ctx.setdefault("标记", set()).add(arg.strip())
            return True
        ok, arg = _is_whole_call(s, "取消标记")
        if ok:
            ctx.setdefault("标记", set()).discard(arg.strip())
            return True
        ok, arg = _is_whole_call(s, "进入管线")
        if ok:
            name = str(arg or "").strip()
            if not name:
                raise 解释器错误("进入管线缺少管线名")
            planned = self._映射规划路名(name, ctx)
            if planned not in self.阶段表 and planned not in self.机制表:
                raise 解释器错误(f"进入管线「{name}」不存在")
            self._执行阶段管线(planned, ctx, 请求名=name)
            return True

        # 护盾副作用（整式）
        ok, arg = _is_whole_call(s, "扣减护盾层")
        if ok:
            self._call_func("扣减护盾层", arg, ctx)
            return True
        ok, arg = _is_whole_call(s, "扣减治疗吸收层")
        if ok:
            self._call_func("扣减治疗吸收层", arg, ctx)
            return True
        ok, arg = _is_whole_call(s, "执行过量转盾")
        if ok:
            self._call_func("执行过量转盾", arg, ctx)
            return True
        ok, arg = _is_whole_call(s, "生成护盾")
        if ok:
            self._call_func("生成护盾", arg, ctx)
            return True

        # 一般表达式（判定 / 终端标识）
        try:
            val = _ExprEval(ctx, self).eval(s)
        except 解释器错误:
            self._unknown_formulas.append(s)
            raise 解释器错误(
                f"无法求值步骤{step.步骤序}「{step.步骤名}」公式：{s[:120]}"
            )

        # 终端标识 / 数值写出：一律成功，走成功跳转
        if _is_terminal_statement(s, val):
            return True
        return bool(val)

    def _do_set(self, fname: str, arg: str, ctx: dict[str, Any]) -> bool:
        if fname == "设结算目标":
            role = arg.strip()
            if role not in ("攻方", "守方"):
                raise 解释器错误(f"设结算目标 非法角色: {role}")
            ctx["结算目标"] = role
            self._sync_settle_life_from_side(ctx)
            return True
        key = _set_key(fname)
        val = _ExprEval(ctx, self).eval(arg)
        # 角色字面量保留字符串
        if key in ("伤害类型", "伤害来源", "治疗来源", "命中位", "暴击位", "效果类型", "叠加规则",
                   "刷新规则", "生效时机", "数值类型", "互斥组", "控制递减组",
                   "效果代号", "控制行为", "叠加动作", "标签",
                   "快照时机", "快照属性列表", "驱散优先级", "驱散类型",
                   "驱散类型等级", "可窃取", "到期动作"):
            ctx[key] = val if isinstance(val, str) else str(val)
        else:
            try:
                ctx[key] = float(val)
            except (TypeError, ValueError):
                ctx[key] = val
        if key == "结算目标生命":
            self._write_settle_life(ctx)
        if key == "当前生命":
            side = self._settle_side(ctx)
            try:
                life = float(ctx[key])
            except (TypeError, ValueError):
                life = 0.0
            side["当前生命"] = life
            side["生命"] = life
            ctx["结算目标生命"] = life
        # 设过量转盾量：有标记「过量转盾开启」时追加全伤害容量盾
        if key == "过量转盾量":
            try:
                amt = float(ctx.get("过量转盾量") or 0.0)
            except (TypeError, ValueError):
                amt = 0.0
            if amt > 1e-9:
                _盾.执行过量转盾(ctx)
        return True

    # -- 函数 / 标识 -----------------------------------------------------------

    def _resolve_ident(self, name: str, ctx: dict[str, Any]) -> Any:
        # 字面量
        if name in 语法.布尔字面量:
            return name == "真"
        extra_lits = frozenset()
        for attr in (
            "快照时机字面量",
            "驱散类型等级字面量",
            "是否字面量",
            "模式字面量",
        ):
            extra_lits |= getattr(语法, attr, frozenset())
        if name in (
            语法.伤害类型字面量
            | 语法.伤害来源字面量
            | 语法.治疗来源字面量
            | 语法.命中位字面量
            | 语法.暴击位字面量
            | 语法.效果类型字面量
            | 语法.叠加规则字面量
            | 语法.刷新规则字面量
            | 语法.生效时机字面量
            | 语法.数值类型字面量
            | 语法.空字面量
            | 语法.角色枚举
            | 语法.标记枚举
            | 语法.状态枚举
            | extra_lits
        ):
            return name
        if name in ctx:
            return ctx[name]
        if name == "护盾层数":
            return float(len(ctx.get("护盾层") or []))
        if name == "治疗吸收层数":
            return float(len(ctx.get("治疗吸收层") or []))
        raise 解释器错误(f"未知标识符「{name}」")

    def _call_func(self, name: str, arg: str, ctx: dict[str, Any]) -> Any:
        arg = arg.strip()
        if name == "随机":
            if arg:
                raise 解释器错误("随机() 无参")
            rng = ctx.get("rng")
            if rng is not None:
                return float(rng.random())
            return 1.0  # 无 RNG：不触发概率

        if name == "最大":
            parts = _split_top_args(arg)
            if len(parts) != 2:
                raise 解释器错误(f"最大() 须两参数：{arg}")
            a = float(_ExprEval(ctx, self).eval(parts[0]))
            b = float(_ExprEval(ctx, self).eval(parts[1]))
            return max(a, b)
        if name == "最小":
            parts = _split_top_args(arg)
            if len(parts) != 2:
                raise 解释器错误(f"最小() 须两参数：{arg}")
            a = float(_ExprEval(ctx, self).eval(parts[0]))
            b = float(_ExprEval(ctx, self).eval(parts[1]))
            return min(a, b)
        if name == "钳制":
            parts = _split_top_args(arg)
            if len(parts) != 3:
                raise 解释器错误(f"钳制() 须三参数：{arg}")
            x = float(_ExprEval(ctx, self).eval(parts[0]))
            lo = float(_ExprEval(ctx, self).eval(parts[1]))
            hi = float(_ExprEval(ctx, self).eval(parts[2]))
            return max(lo, min(hi, x))

        if name == "有标记":
            return arg.strip() in (ctx.get("标记") or set())
        if name == "没有标记":
            return arg.strip() not in (ctx.get("标记") or set())
        if name == "没有状态":
            return arg.strip() not in _side_states(ctx.get("守方"))
        if name == "结算目标有标记":
            side = self._settle_side(ctx)
            flag = arg.strip()
            if flag in _side_flags(side):
                return True
            # 管线标记（已死亡/攻方已死亡等）与结算目标共用
            return flag in (ctx.get("标记") or set())

        if name == "有状态":
            # 默认看守方（受击方）
            return arg.strip() in _side_states(ctx.get("守方"))
        if name == "结算目标有状态":
            return arg.strip() in _side_states(self._settle_side(ctx))

        if name == "效果":
            parts = _split_top_args(arg)
            if len(parts) not in (1, 2):
                raise 解释器错误(f"效果() 参数错误：{arg}")
            attr = parts[0].strip()
            role = parts[1].strip() if len(parts) == 2 else None
            return self._lookup_effect(attr, role, ctx)

        if name in ("攻方", "守方"):
            panel = ctx.get(name) or {}
            return _panel_get(panel, arg.strip(), 0.0)

        if name == "结算目标属性":
            side = self._settle_side(ctx)
            key = arg.strip()
            if key in ("生命值", "生命上限"):
                return float(side.get("生命上限") or side.get("生命值") or side.get("生命") or 0.0)
            if key in ("生命", "当前生命"):
                return float(ctx.get("结算目标生命") or side.get("生命") or 0.0)
            return _panel_get(side, key, 0.0)

        if name == "取护盾层":
            idx = int(float(_ExprEval(ctx, self).eval(arg)))
            return _盾.取护盾层(ctx.get("护盾层"), idx)

        if name == "护盾剩余容量":
            sh = _ExprEval(ctx, self).eval(arg)
            return float(_盾.护盾剩余容量(sh if isinstance(sh, dict) else None))

        if name == "护盾类型可吸收":
            parts = _split_top_args(arg)
            if len(parts) != 2:
                raise 解释器错误("护盾类型可吸收 须两参数")
            sh = _ExprEval(ctx, self).eval(parts[0])
            dmg_t = _ExprEval(ctx, self).eval(parts[1])
            return bool(
                _盾.护盾类型可吸收(sh if isinstance(sh, dict) else None, dmg_t)
            )

        if name == "扣减护盾层":
            parts = _split_top_args(arg)
            if len(parts) != 2:
                raise 解释器错误("扣减护盾层 须两参数")
            sh = _ExprEval(ctx, self).eval(parts[0])
            amt = float(_ExprEval(ctx, self).eval(parts[1]))
            if isinstance(sh, dict):
                was_alive = _盾.护盾仍存活(sh)
                _盾.扣减护盾层(sh, amt)
                if was_alive and not _盾.护盾仍存活(sh):
                    ctx["本次破盾数"] = float(ctx.get("本次破盾数") or 0.0) + 1.0
                    # 保持列表一致：不在中途删除（表用游标扫），写回时压实存活层
                    if _盾.护盾需破盾结束Buff(sh):
                        _盾.执行破盾结束Buff(ctx, [sh])
            return True

        if name == "执行过量转盾":
            return float(_盾.执行过量转盾(ctx))

        if name == "生成护盾":
            # 生成护盾(容量) 或 生成护盾(容量,吸收类型)
            parts = _split_top_args(arg) if arg else []
            if not parts:
                raise 解释器错误("生成护盾 至少须容量参数")
            cap = float(_ExprEval(ctx, self).eval(parts[0]))
            atype = "全伤害"
            if len(parts) >= 2:
                atype = str(_ExprEval(ctx, self).eval(parts[1]))
            sh = _盾.新建容量盾(cap, 吸收类型=atype)
            layers = ctx.get("护盾层")
            if not isinstance(layers, list):
                layers = []
                ctx["护盾层"] = layers
            layers.append(sh)
            ctx["护盾层数"] = float(len(layers))
            side = self._settle_side(ctx)
            if side.get("护盾层") is not layers:
                side["护盾层"] = layers
            return True

        if name == "取治疗吸收层":
            idx = int(float(_ExprEval(ctx, self).eval(arg)))
            return _盾.取治疗吸收层(ctx.get("治疗吸收层"), idx)

        if name == "治疗吸收剩余容量":
            sh = _ExprEval(ctx, self).eval(arg)
            if isinstance(sh, dict):
                return float(_盾._heal_absorb_cap(sh))
            return 0.0

        if name == "扣减治疗吸收层":
            parts = _split_top_args(arg)
            if len(parts) != 2:
                raise 解释器错误("扣减治疗吸收层 须两参数")
            sh = _ExprEval(ctx, self).eval(parts[0])
            amt = float(_ExprEval(ctx, self).eval(parts[1]))
            if isinstance(sh, dict):
                _盾.扣减治疗吸收层(sh, amt)
            return True

        if name == "生成治疗吸收":
            parts = _split_top_args(arg) if arg else []
            if not parts:
                raise 解释器错误("生成治疗吸收 至少须容量参数")
            cap = float(_ExprEval(ctx, self).eval(parts[0]))
            sh = _盾.新建治疗吸收层(cap)
            layers = ctx.get("治疗吸收层")
            if not isinstance(layers, list):
                layers = []
                ctx["治疗吸收层"] = layers
            layers.append(sh)
            ctx["治疗吸收层数"] = float(len(layers))
            side = self._settle_side(ctx)
            if side.get("治疗吸收层") is not layers:
                side["治疗吸收层"] = layers
            return True

        # 异常 / 光环 / 驱散助手 → 内核.光环运行时
        from 战斗模拟.内核 import 光环运行时 as _光环

        if name == "效果标签被免疫":
            # 表侧免疫标签尚未全量接线；缺省不免疫
            return False
        if name == "结算目标递减系数":
            return float(ctx.get("递减系数") or 1.0)

        if name == "互斥组冲突":
            group = _ExprEval(ctx, self).eval(arg) if arg else ctx.get("互斥组")
            side = self._settle_side(ctx)
            return bool(_光环.互斥组冲突(side, group))

        if name == "有同效果":
            code = _ExprEval(ctx, self).eval(arg) if arg else ctx.get("效果代号")
            side = self._settle_side(ctx)
            hit = _光环.有同效果(side, str(code or ""))
            if hit:
                found = _光环.查找同效果(side, str(code or ""))
                if found:
                    ctx["当前层数"] = float(found[0].层数)
                    ctx["剩余时长"] = float(found[0].剩余时长)
                    ctx["旧强度"] = float(ctx.get("旧强度") or found[0].层数)
            return hit

        if name == "移除互斥组效果":
            group = _ExprEval(ctx, self).eval(arg) if arg else ctx.get("互斥组")
            side = self._settle_side(ctx)
            removed = _光环.移除互斥组效果(side, group)
            self._cancel_aura_events(ctx, removed)
            return True

        if name == "移除同效果":
            code = _ExprEval(ctx, self).eval(arg) if arg else ctx.get("效果代号")
            side = self._settle_side(ctx)
            removed = _光环.移除同效果(side, str(code or ""))
            self._cancel_aura_events(ctx, removed)
            return True

        if name == "移除效果":
            code = _ExprEval(ctx, self).eval(arg) if arg else ctx.get("效果代号")
            side = self._settle_side(ctx)
            uid = ctx.get("_当前光环uid")
            removed = _光环.移除效果(side, str(code or ""), uid=uid)
            self._cancel_aura_events(ctx, removed)
            return True

        if name == "挂载效果":
            parts = _split_top_args(arg)
            if len(parts) != 3:
                raise 解释器错误("挂载效果 须三参数 (效果代号,层数,剩余时长)")
            ev = _ExprEval(ctx, self)
            code = ev.eval(parts[0])
            stacks = float(ev.eval(parts[1]))
            dur = float(ev.eval(parts[2]))
            side = self._settle_side(ctx)
            inst = _光环.挂载效果(
                side,
                str(code),
                stacks,
                dur,
                ctx=ctx,
            )
            ctx["异常挂载结果"] = 1.0
            ctx["异常层数"] = float(inst.层数)
            ctx["当前层数"] = float(inst.层数)
            ctx["剩余时长"] = float(inst.剩余时长)
            ctx["_当前光环uid"] = inst.uid
            ctx["_当前光环"] = inst
            # 若引擎挂在上下文，立即调度到期/衰减/跳动
            eng = ctx.get("_战斗引擎")
            if eng is not None and hasattr(eng, "调度光环存续"):
                try:
                    eng.调度光环存续(
                        目标侧=side,
                        光环=inst,
                        来源id=str((ctx.get("攻方") or {}).get("id") or ""),
                        目标id=str(side.get("id") or ""),
                    )
                except Exception:
                    pass
            return True

        if name == "写入快照":
            attr_list = _ExprEval(ctx, self).eval(arg) if arg else ctx.get("快照属性列表")
            aura = ctx.get("_当前光环")
            if aura is None:
                # 回落：按效果代号找守方上最新实例
                side = self._settle_side(ctx)
                found = _光环.查找同效果(side, str(ctx.get("效果代号") or ""))
                aura = found[-1] if found else None
            if aura is None:
                return True
            _光环.写入快照(aura, ctx.get("攻方"), attr_list if attr_list not in (None, "") else ctx.get("快照属性列表"))
            ctx["_当前光环"] = aura
            return True

        if name == "有可驱散效果":
            side = self._settle_side(ctx)
            tried = ctx.setdefault("_驱散已尝试uids", set())
            return bool(_光环.有可驱散效果(side, 已尝试uids=tried))

        if name == "选取可驱散效果":
            # 参数为驱散优先级基准（字符串/数值）；排序已在运行时按光环.驱散优先级
            side = self._settle_side(ctx)
            tried = ctx.setdefault("_驱散已尝试uids", set())
            aura = _光环.选取可驱散效果(side, 已尝试uids=tried)
            if aura is None:
                return False
            tried.add(aura.uid)
            _光环.填充上下文自光环(ctx, aura)
            return True

        if name == "匹配驱散类型":
            parts = _split_top_args(arg)
            if len(parts) != 2:
                raise 解释器错误("匹配驱散类型 须两参数 (标签,驱散类型)")
            ev = _ExprEval(ctx, self)
            tag = ev.eval(parts[0])
            dtype = ev.eval(parts[1])
            return bool(_光环.匹配驱散类型(tag, dtype))

        if name == "转移效果":
            parts = _split_top_args(arg)
            if len(parts) != 2:
                raise 解释器错误("转移效果 须两参数 (效果代号,攻方|守方)")
            ev = _ExprEval(ctx, self)
            code = str(ev.eval(parts[0]) or "")
            dest_role = str(parts[1]).strip()
            # 第二参数常为字面 攻方/守方
            if dest_role not in ("攻方", "守方"):
                dest_role = str(ev.eval(parts[1]) or "").strip()
            src = self._settle_side(ctx)
            dst = ctx.get(dest_role)
            if not isinstance(dst, dict):
                dst = {}
                ctx[dest_role] = dst
            uid = ctx.get("_当前光环uid")
            inst = _光环.转移效果(src, dst, code, uid=uid)
            if inst is None:
                return False
            ctx["_当前光环"] = inst
            ctx["_当前光环uid"] = inst.uid
            return True

        if name == "执行到期动作":
            # 复杂到期动作（再挂载/爆炸伤）后续扩展；此处记标记并返回成功
            action = _ExprEval(ctx, self).eval(arg) if arg else ctx.get("到期动作")
            ctx["_最近到期动作"] = action
            marks = ctx.setdefault("标记", set())
            if isinstance(marks, set):
                marks.add("已执行到期动作")
            return True

        raise 解释器错误(f"未知函数「{name}(...)」")

    def _lookup_effect(self, attr: str, role: str | None, ctx: dict[str, Any]) -> float:
        over = ctx.get("效果覆盖") or {}
        if attr in over:
            return float(over[attr])
        mode = self._当前模式(ctx)
        candidates = [attr]
        if "(" not in attr:
            candidates.append(f"{attr}({mode})")
            other = "PVP" if mode == "PVE" else "PVE"
            candidates.append(f"{attr}({other})")
        for name in candidates:
            if name in over:
                return float(over[name])

        def _has(panel: dict | None, key: str) -> bool:
            if not panel:
                return False
            if key in panel and panel[key] is not None:
                return True
            nested = panel.get("面板") if isinstance(panel.get("面板"), dict) else None
            return bool(nested and key in nested and nested[key] is not None)

        atk = ctx.get("攻方") or {}
        dfd = ctx.get("守方") or {}
        if role == "攻方":
            for name in candidates:
                if _has(atk, name):
                    return _panel_get(atk, name, 0.0)
            return 0.0
        if role == "守方":
            for name in candidates:
                if _has(dfd, name):
                    return _panel_get(dfd, name, 0.0)
            return 0.0
        prefer_def = any(x in attr for x in ("免伤", "抗性"))
        order = (dfd, atk) if prefer_def else (atk, dfd)
        for panel in order:
            for name in candidates:
                if _has(panel, name):
                    return _panel_get(panel, name, 0.0)
        if "效果" in attr and not any(
            x in attr for x in ("闪避", "格挡", "暴击", "附着", "控制")
        ):
            if any(x in attr for x in ("免伤",)):
                return 0.0
            return 1.0
        return 0.0

    def _settle_side(self, ctx: dict[str, Any]) -> dict[str, Any]:
        """返回结算目标面板的真实引用（非拷贝），供状态/护盾读写。"""
        role = ctx.get("结算目标") or "守方"
        side = ctx.get(role)
        if not isinstance(side, dict):
            side = {}
            ctx[role] = side
        return side

    def _sync_settle_life_from_side(self, ctx: dict[str, Any]) -> None:
        """切换结算目标时：装载该侧生命，并绑定该侧护盾层到 ctx（反伤打攻方盾关键）。"""
        side = self._settle_side(ctx)
        life = float(side.get("生命") or side.get("当前生命") or 0.0)
        ctx["结算目标生命"] = life
        ctx["当前生命"] = life
        if not isinstance(side.get("护盾层"), list):
            side["护盾层"] = list(side.get("护盾层") or [])
        # 同一 list 引用：FIFO 就地扣减直接写到面板
        ctx["护盾层"] = side["护盾层"]
        ctx["护盾层数"] = float(len(ctx["护盾层"]))
        ctx["护盾游标"] = float(ctx.get("护盾游标") or 0.0)
        if not isinstance(side.get("治疗吸收层"), list):
            side["治疗吸收层"] = list(side.get("治疗吸收层") or [])
        ctx["治疗吸收层"] = side["治疗吸收层"]
        ctx["治疗吸收层数"] = float(len(ctx["治疗吸收层"]))
        ctx["治疗吸收游标"] = float(ctx.get("治疗吸收游标") or 0.0)

    def _write_settle_life(self, ctx: dict[str, Any]) -> None:
        role = ctx.get("结算目标") or "守方"
        side = ctx.get(role)
        if not isinstance(side, dict):
            return
        life = float(ctx.get("结算目标生命") or 0.0)
        side["生命"] = life
        side["当前生命"] = life
        # 压实存活护盾写回
        alive = _盾.存活护盾列表(ctx.get("护盾层") or side.get("护盾层") or [])
        side["护盾层"] = alive
        ctx["护盾层"] = alive
        ctx["护盾层数"] = float(len(alive))
        ha = _盾.存活治疗吸收列表(ctx.get("治疗吸收层") or side.get("治疗吸收层") or [])
        side["治疗吸收层"] = ha
        ctx["治疗吸收层"] = ha
        ctx["治疗吸收层数"] = float(len(ha))
        # 若面板挂了实体引用则同步
        ent = side.get("_实体") or side.get("实体")
        if ent is not None:
            try:
                ent.生命 = life
                if hasattr(ent, "护盾层"):
                    ent.护盾层 = list(alive)
                if life <= 0 and hasattr(ent, "存活"):
                    ent.存活 = False
            except Exception:
                pass

    def _写回生命(self, ctx: dict[str, Any]) -> None:
        # 两侧都写回（吸血改攻方、扣血改守方）；以各侧面板生命/护盾为准
        for role in ("攻方", "守方"):
            side = ctx.get(role)
            if not isinstance(side, dict):
                continue
            if "生命" in side or "当前生命" in side:
                life = float(side.get("生命") or side.get("当前生命") or 0.0)
                side["生命"] = life
                side["当前生命"] = life
            if isinstance(side.get("护盾层"), list):
                side["护盾层"] = _盾.存活护盾列表(side["护盾层"])
            ent = side.get("_实体") or side.get("实体")
            if ent is not None:
                try:
                    ent.生命 = float(side.get("生命") or 0.0)
                    if hasattr(ent, "护盾层"):
                        ent.护盾层 = list(side.get("护盾层") or [])
                    if float(side.get("生命") or 0.0) <= 0 and hasattr(ent, "存活"):
                        ent.存活 = False
                except Exception:
                    pass
        # 当前结算目标再刷一次（含压实）
        self._write_settle_life(ctx)


    def _cancel_aura_events(self, ctx: dict[str, Any], auras: list) -> None:
        """移除光环时取消已调度的到期/衰减/跳动事件。"""
        eng = ctx.get("_战斗引擎")
        sch = getattr(eng, "调度器", None) if eng is not None else ctx.get("_事件调度器")
        if sch is None:
            return
        from 战斗模拟.内核 import 光环运行时 as _光环

        for a in auras or []:
            try:
                _光环.取消光环调度(sch, a)
            except Exception:
                pass

    def opcode_coverage(self) -> dict[str, Any]:
        """统计已加载管线中公式模式覆盖概况。"""
        total = 0
        covered = 0
        patterns = {
            "设*": 0,
            "标记/取消": 0,
            "判定/比较": 0,
            "终端": 0,
            "护盾": 0,
            "其他": 0,
        }
        for steps in self.管线表.values():
            for st in steps:
                total += 1
                s = _compact(st.判定公式)
                hit = False
                for fn in _SET_FUNCS:
                    if s.startswith(fn + "("):
                        patterns["设*"] += 1
                        hit = True
                        break
                if hit:
                    covered += 1
                    continue
                if s.startswith("标记(") or s.startswith("取消标记("):
                    patterns["标记/取消"] += 1
                    covered += 1
                    continue
                if any(
                    s.startswith(p)
                    for p in (
                        "护盾类型可吸收",
                        "扣减护盾层",
                        "取护盾层",
                        "护盾剩余容量",
                    )
                ) or "护盾" in s[:6]:
                    patterns["护盾"] += 1
                    covered += 1
                    continue
                if s in (
                    "最终伤害",
                    "实际治疗",
                    "结算伤害",
                    "实际扣血",
                    "异常挂载结果",
                    "治疗值",
                ):
                    patterns["终端"] += 1
                    covered += 1
                    continue
                if any(op in s for op in ("==", "!=", "<", ">", "有标记", "有状态", "随机")):
                    patterns["判定/比较"] += 1
                    covered += 1
                    continue
                patterns["其他"] += 1
                covered += 1  # evaluator attempts all
        pct = (100.0 * covered / total) if total else 0.0
        return {
            "总步骤": total,
            "可解析步骤": covered,
            "覆盖率%": round(pct, 1),
            "模式": patterns,
            "管线": self.管线名列表(),
        }


def _is_terminal_statement(s: str, val: Any) -> bool:
    """整式为裸标识/数值写出时视为成功。"""
    if s in 语法.管线局部标识符:
        return True
    if isinstance(val, (int, float)) and not isinstance(val, bool):
        # 纯数值表达式作为步骤体（少见）— 仍当成功
        if not any(op in s for op in ("==", "!=", "<=", ">=", "<", ">", "||", "&&")):
            if not any(
                s.startswith(fn + "(")
                for fn in (
                    "有标记",
                    "有状态",
                    "结算目标有标记",
                    "结算目标有状态",
                    "随机",
                    "护盾类型可吸收",
                    "互斥组冲突",
                    "有同效果",
                    "效果标签被免疫",
                )
            ):
                # 形如 `结算伤害` 已覆盖；`实际治疗` 同
                if re.fullmatch(
                    r"[\u4e00-\u9fffA-Za-z_][\u4e00-\u9fffA-Za-z0-9_%]*", s
                ):
                    return True
    return False


FlowInterpreter = 流程解释器
PipeResult = 管线执行结果
build_context = 构建上下文
