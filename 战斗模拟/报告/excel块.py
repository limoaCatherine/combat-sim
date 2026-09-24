"""将结果写成清晰块 + 新鲜度时间戳。"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook


def 写结果块(
    path: str | Path,
    *,
    块标题: str,
    行数据: list[list[Any]],
    表名: str = "结果",
) -> Path:
    p = Path(path)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = 表名
    ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    ws.cell(1, 1, 块标题)
    ws.cell(2, 1, f"生成时间: {ts}")
    for i, row in enumerate(行数据, start=4):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return p


write_excel_blocks = 写结果块
