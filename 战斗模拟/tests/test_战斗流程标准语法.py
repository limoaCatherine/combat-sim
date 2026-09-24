# -*- coding: utf-8 -*-
"""战斗流程表 — 中文标准 DSL 回归（SimC Phase A · 6 管线（效果事件含驱散））。"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from 战斗模拟.load.战斗流程 import 加载战斗流程, 战斗流程加载错误
from 战斗模拟.领域.流程语法 import (
    从属性总表加载用途名,
    校验判定公式,
    校验步骤名,
    规范化步骤名,
    流程语法错误,
    角色枚举,
    管线名枚举,
    规范化管线名,
    管线名别名,
)

FW = Path("/workspace/combat-framework/战斗数值框架.xlsx")

EXPECTED_PIPES = [
    "伤害PVE",
    "伤害PVP",
    "治疗PVE",
    "治疗PVP",
    "效果事件PVE",
    "效果事件PVP",
]
FORBIDDEN_PIPES = (
    "扣血PVE",
    "扣血PVP",
    "回血PVE",
    "回血PVP",
    "吸血反伤PVE",
    "吸血反伤PVP",
    "伤害仇恨PVE",
    "治疗仇恨PVE",
    "命中后结算PVE",
    "命中后结算PVP",
    "治疗后结算PVE",
    "治疗后结算PVP",
    "驱散PVE",
    "驱散PVP",
)

_EN_BANNED = (
    "setPipelineDamage",
    "formulaParam",
    "pipelineDamage",
    "isPhysicalDamage",
    "isMagicalDamage",
    "setDodged",
    "setCritical",
    "casterAttr",
    "targetAttr",
    "damageContext",
)

扣血段关键步骤 = [
    "判定结算目标免疫或无敌",
    "初始化护盾游标",
    "判定还有护盾可吸收",
    "取当前层护盾",
    "判定本层可吸收",
    "计算本层吸收",
    "扣减本层护盾",
    "护盾游标前进",
    "执行扣生命",
    "判定锁1血",
    "应用锁1血",
    "记录实际扣血",
    "扣血段收束",
]


@pytest.fixture(scope="module")
def flow_wb():
    if not FW.is_file():
        pytest.skip("框架表不在 box")
    wb = load_workbook(FW, data_only=True)
    yield wb
    wb.close()


@pytest.fixture(scope="module")
def attr_names(flow_wb):
    return 从属性总表加载用途名(flow_wb)


@pytest.fixture(scope="module")
def pipelines(flow_wb):
    return 加载战斗流程(flow_wb)


def test_pipelines_present(pipelines):
    assert set(pipelines) == set(EXPECTED_PIPES)
    assert len(pipelines) == 6
    assert set(EXPECTED_PIPES) <= 管线名枚举
    for name in EXPECTED_PIPES:
        assert name in pipelines
        assert pipelines[name], f"管线 {name} 无步骤"
    for name in FORBIDDEN_PIPES:
        assert name not in pipelines
    # 短期兼容：旧名异常*/驱散* → 效果事件*
    assert 规范化管线名("异常PVE") == "效果事件PVE"
    assert 规范化管线名("异常PVP") == "效果事件PVP"
    assert 规范化管线名("驱散PVE") == "效果事件PVE"
    assert 规范化管线名("驱散PVP") == "效果事件PVP"
    assert 管线名别名["异常PVE"] == "效果事件PVE"
    assert 管线名别名["驱散PVE"] == "效果事件PVE"
    assert "异常PVE" not in pipelines  # 表侧已切新名
    assert "驱散PVE" not in pipelines  # 驱散已并入效果事件


def test_no_placeholder_formulas(pipelines):
    """禁止步骤整式仅为 1 / 真。"""
    for name, steps in pipelines.items():
        for s in steps:
            f = str(s.判定公式).strip()
            assert f not in {"1", "真", "假"}, (name, s.步骤序, f)


def test_no_调用结算_in_sheet(pipelines):
    for name, steps in pipelines.items():
        for s in steps:
            assert "调用结算" not in s.判定公式, (name, s.步骤序)


def test_every_formula_passes_validator(pipelines, attr_names):
    assert attr_names, "属性总表用途名为空"
    for name, steps in pipelines.items():
        for s in steps:
            try:
                校验判定公式(s.判定公式, attr_names)
            except 流程语法错误 as e:
                pytest.fail(f"{name}#{s.步骤序} {s.步骤名}: {e}")


def test_every_attr_ref_in_master(pipelines, attr_names):
    from 战斗模拟.领域 import 流程语法 as syn

    for name, steps in pipelines.items():
        for s in steps:
            for fn in ("效果", "攻方", "守方", "结算目标属性"):
                for arg, _a, _b in syn._iter_func_calls(s.判定公式, fn):
                    if fn == "效果":
                        parts = syn._split_top_args(arg)
                        assert parts[0] in attr_names, (
                            f"{name}#{s.步骤序} 效果({parts[0]}) ∉ 属性总表"
                        )
                        if len(parts) == 2:
                            assert parts[1] in 角色枚举, (
                                f"{name}#{s.步骤序} 效果 role={parts[1]!r}"
                            )
                    else:
                        assert arg in attr_names, (
                            f"{name}#{s.步骤序} {fn}({arg}) ∉ 属性总表"
                        )


def test_no_english_aviator_left(pipelines):
    for name, steps in pipelines.items():
        for s in steps:
            f = s.判定公式
            for ban in _EN_BANNED:
                assert ban not in f, f"{name}#{s.步骤序} still has {ban}"


def test_step_names_normalized(pipelines):
    for name, steps in pipelines.items():
        for s in steps:
            assert s.步骤名 in 规范化步骤名, (name, s.步骤序, s.步骤名)
            校验步骤名(s.步骤名)


def test_jumps_allow_minus_one(pipelines):
    for name, steps in pipelines.items():
        orders = {s.步骤序 for s in steps}
        allowed = {-1, 0} | orders
        for s in steps:
            assert s.成功跳转 in allowed, (name, s.步骤序, s.成功跳转)
            assert s.失败跳转 in allowed, (name, s.步骤序, s.失败跳转)


def test_damage_总分总_structure(pipelines, attr_names):
    """伤害*：init hygiene → 持续/必中 → 算量 → 夹断 → 内联扣血 → [PVE仇恨] → 写出。"""
    pve = pipelines["伤害PVE"]
    assert len(pve) >= 80
    assert pve[0].判定公式 == "设管线伤害(基础伤害)"
    names = [s.步骤名 for s in pve]
    assert "清除闪避标记" in names
    assert "判定持续伤害来源" in names
    assert "【持续】跳过闪避判定" in names
    assert any(s.判定公式 == "伤害来源==持续" for s in pve)
    assert "判定必中" in names
    assert any("有标记(必中)||命中位==必中" == s.判定公式 for s in pve)
    dodge = next(s for s in pve if s.步骤名 == "判定闪避")
    assert "无敌" not in dodge.判定公式
    assert "判定锁1血" in names
    assert "应用锁1血" in names
    assert "无敌清零吸血量" in names
    assert "应用物理免伤与物理伤害效果" in names
    assert "应用魔法免伤与魔法伤害效果" in names
    assert "最终伤害夹断非负" in names
    assert any(s.判定公式 == "设管线伤害(最大(0,管线伤害))" for s in pve)
    assert any("设吸血量(最终伤害*效果(吸血效果(PVE),攻方))" == s.判定公式 for s in pve)
    for req in 扣血段关键步骤:
        assert req in names, req
    assert pve[-1].步骤名 == "写出伤害管线结果"
    assert pve[-1].判定公式 == "最终伤害"
    assert pve[-1].成功跳转 == -1
    assert any(s.步骤名 == "写出伤害仇恨" for s in pve)
    assert any("设伤害仇恨((实际扣血+反伤量)*仇恨系数)" == s.判定公式 for s in pve)
    assert any(s.判定公式 == "设伤害仇恨(0)" for s in pve)

    pvp = pipelines["伤害PVP"]
    assert len(pvp) >= 80
    assert any("闪避效果(PVP)" in s.判定公式 for s in pvp)
    assert pvp[-1].步骤名 == "写出伤害管线结果"
    assert all(s.步骤名 not in ("写出伤害仇恨", "清零伤害仇恨") for s in pvp)
    for s in list(pve) + list(pvp):
        校验判定公式(s.判定公式, attr_names)


def test_heal_总分总_structure(pipelines, attr_names):
    for mode, name in (("PVE", "治疗PVE"), ("PVP", "治疗PVP")):
        steps = pipelines[name]
        assert steps[0].判定公式 == "设治疗值(基础治疗)"
        assert steps[0].步骤名 == "初始化治疗"
        assert any(
            s.判定公式 == f"设治疗值(治疗值*效果(施法治疗效果({mode}),攻方))"
            for s in steps
        )
        assert any("暴击位==不可暴击||有标记(不可暴击)" == s.判定公式 for s in steps)
        assert any(s.判定公式 == "设治疗值(最大(0,治疗值))" for s in steps)
        assert any(s.步骤名 == "判定零治疗短路" for s in steps)
        assert any(s.步骤名 == "判定还有治疗吸收" for s in steps)
        assert any(s.步骤名 == "判定过量转盾" for s in steps)
        assert any(s.步骤名 == "过量转化护盾" for s in steps)
        assert any(s.步骤名 == "回血收束" for s in steps)
        assert steps[-1].步骤名 == "写出治疗管线结果"
        assert steps[-1].判定公式 == "实际治疗"
        assert steps[-1].成功跳转 == -1
        for s in steps:
            校验判定公式(s.判定公式, attr_names)

    pve = pipelines["治疗PVE"]
    assert any(s.判定公式 == "设治疗仇恨(实际治疗*0.05*仇恨系数)" for s in pve)
    assert any(s.判定公式 == "设治疗仇恨(0)" for s in pve)
    assert any(s.判定公式 == "实际治疗>0" for s in pve)
    pvp = pipelines["治疗PVP"]
    assert all(s.步骤名 not in ("写出治疗仇恨", "清零治疗仇恨") for s in pvp)


def test_effect_event_structure(pipelines, attr_names):
    """效果事件*：拒绝门 → 叠层 → 挂载 → 快照 → handoff → 层衰减 → 到期 + 驱散/偷取段。"""
    for mode, name in (("PVE", "效果事件PVE"), ("PVP", "效果事件PVP")):
        steps = pipelines[name]
        assert 80 <= len(steps) <= 100
        names = [s.步骤名 for s in steps]
        assert steps[0].判定公式 == "设异常挂载结果(0)"
        assert any("控制命中" in s.步骤名 for s in steps)
        assert any(
            f"效果(控制效果({mode}))" in s.判定公式 for s in steps
        )
        assert any("叠层" in s.步骤名 for s in steps)
        assert any("挂载效果(" in s.判定公式 for s in steps)
        assert any(s.步骤名 == "【快照】判定施加时快照" for s in steps)
        assert any(s.判定公式 == "写入快照(快照属性列表)" for s in steps)
        assert any(s.步骤名 == "写出跳伤请求" for s in steps)
        assert any(s.步骤名 == "写出跳疗请求" for s in steps)
        assert any(s.判定公式 == "设伤害来源(持续)" for s in steps)
        assert any(s.判定公式 == "设治疗来源(持续)" for s in steps)
        assert any(s.步骤名 == "【层衰减】判定层衰减入口" for s in steps)
        assert any("层衰减间隔>0" in s.判定公式 for s in steps)
        assert any("设异常层数(最大(0,当前层数-层衰减数量))" == s.判定公式 for s in steps)
        assert any(s.步骤名 == "【到期】判定到期入口" for s in steps)
        assert any(s.判定公式 == "有标记(异常到期结算)" for s in steps)
        assert any(s.步骤名 == "【到期】执行到期动作" for s in steps)
        # 驱散/偷取段并入末尾
        assert any(s.判定公式 == "设驱散结果(0)" for s in steps)
        assert any(s.判定公式 == "有可驱散效果()" for s in steps)
        assert any("转移效果(效果代号,攻方)" == s.判定公式 for s in steps)
        assert any(s.步骤名 == "驱散结果" for s in steps)
        assert "调用结算" not in "".join(s.判定公式 for s in steps)
        for s in steps:
            校验判定公式(s.判定公式, attr_names)
            assert str(s.判定公式).strip() not in {"1", "真"}


def test_dispel_structure(pipelines, attr_names):
    """驱散段已并入效果事件* 末尾：预处理 → 选取/匹配 → 移除或转移 → 驱散结果终端。"""
    for name in ("效果事件PVE", "效果事件PVP"):
        steps = pipelines[name]
        # 从「初始化驱散结果」起截到终端（并入效果事件末尾的独立入口段）
        idx = next(i for i, s in enumerate(steps) if s.判定公式 == "设驱散结果(0)")
        seg = steps[idx:]
        assert 25 <= len(seg) <= 35
        assert seg[0].判定公式 == "设驱散结果(0)"
        assert any(s.判定公式 == "设已驱散数(0)" for s in seg)
        assert any("有可驱散效果()" == s.判定公式 for s in seg)
        assert any("选取可驱散效果(驱散优先级)" == s.判定公式 for s in seg)
        assert any("匹配驱散类型(标签,驱散类型)" == s.判定公式 for s in seg)
        assert any("移除效果(效果代号)" == s.判定公式 for s in seg)
        assert any("转移效果(效果代号,攻方)" == s.判定公式 for s in seg)
        assert seg[-1].步骤名 == "驱散结果"
        assert seg[-1].判定公式 == "驱散结果"
        assert seg[-1].成功跳转 == -1
        assert "调用结算" not in "".join(s.判定公式 for s in seg)
        for s in seg:
            校验判定公式(s.判定公式, attr_names)
            assert str(s.判定公式).strip() not in {"1", "真", "假"}


def test_heal_crit_attrs_in_master(attr_names):
    for n in (
        "治疗暴击效果(PVE)",
        "治疗暴击效果(PVP)",
        "治疗暴击伤害效果(PVE)",
        "治疗暴击伤害效果(PVP)",
        "吸血效果(PVE)",
        "吸血效果(PVP)",
        "反伤效果(PVE)",
        "反伤效果(PVP)",
        "控制效果(PVE)",
        "控制效果(PVP)",
    ):
        assert n in attr_names


def test_lifesteal_effect_binding(pipelines):
    for name in ("伤害PVE", "伤害PVP"):
        steps = pipelines[name]
        assert not any("设对应吸血" in s.判定公式 for s in steps)
        assert not any("判定物理通道" == s.步骤名 for s in steps)
        assert any("效果(吸血效果(" in s.判定公式 and ",攻方)" in s.判定公式 for s in steps)
        assert any("效果(反伤效果(" in s.判定公式 and ",守方)" in s.判定公式 for s in steps)


def test_accept_总分总_dsl(attr_names):
    校验判定公式("设最终伤害(管线伤害)", attr_names)
    校验判定公式("取消标记(已闪避)", attr_names)
    校验判定公式("伤害来源==持续", attr_names)
    校验判定公式("有标记(必中)||命中位==必中", attr_names)
    校验判定公式("设管线伤害(最大(0,管线伤害))", attr_names)
    校验判定公式("设治疗值(基础治疗)", attr_names)
    校验判定公式("暴击位==不可暴击||有标记(不可暴击)", attr_names)
    校验判定公式("设异常挂载结果(0)", attr_names)
    校验判定公式("互斥组冲突(互斥组)", attr_names)
    校验判定公式("挂载效果(效果代号,异常层数,剩余时长)", attr_names)
    校验判定公式("随机()<钳制(效果(控制效果(PVE)),0,1)", attr_names)
    校验判定公式("设跳伤伤害(异常跳伤基础)", attr_names)
    校验判定公式("设伤害来源(持续)", attr_names)
    校验判定公式("设治疗来源(持续)", attr_names)
    校验判定公式("有标记(异常到期结算)", attr_names)
    校验判定公式("写入快照(快照属性列表)", attr_names)
    校验判定公式("快照时机==施加时", attr_names)
    校验判定公式("有标记(层衰减结算)", attr_names)
    校验判定公式("设驱散结果(0)", attr_names)
    校验判定公式("有可驱散效果()", attr_names)
    校验判定公式("匹配驱散类型(标签,驱散类型)", attr_names)
    校验判定公式("转移效果(效果代号,攻方)", attr_names)
    校验判定公式("真", attr_names)
    with pytest.raises(流程语法错误, match="调用结算"):
        校验判定公式("调用结算(扣血,最终伤害)", attr_names)


def test_reject_unknown_attr():
    with pytest.raises(流程语法错误, match="不在属性总表"):
        校验判定公式("效果(不存在的用途名XYZ)", {"闪避效果(PVE)"})


def test_reject_bad_effect_role():
    with pytest.raises(流程语法错误, match="第二参数"):
        校验判定公式("效果(闪避效果(PVE),中立)", {"闪避效果(PVE)"})


def test_reject_english_aviator():
    with pytest.raises(流程语法错误):
        校验判定公式("setPipelineDamage(damageContext.baseDamage)", set())
    with pytest.raises(流程语法错误):
        校验判定公式("random() < formulaParam('DodgeEffectPVE')", set())


def test_reject_bad_status_and_shield_kind():
    with pytest.raises(流程语法错误, match="有状态"):
        校验判定公式("有状态(沉默)", set())
    with pytest.raises(流程语法错误, match="护盾吸收"):
        校验判定公式("护盾吸收(结算伤害,火焰)", set())
    with pytest.raises(流程语法错误, match="护盾类型可吸收"):
        校验判定公式("护盾类型可吸收(当前护盾,火焰)", set())


def test_loader_accepts_minus_one():
    wb = Workbook()
    ws = wb.active
    ws.title = "战斗流程"
    ws.cell(1, 1, "战斗流程")
    ws.cell(3, 1, "伤害PVE")
    for i, h in enumerate(("管线", "步骤序", "步骤名", "判定公式", "成功跳转", "失败跳转")):
        ws.cell(4, 1 + i, h)
    ws.cell(5, 1, "伤害PVE")
    ws.cell(5, 2, 1)
    ws.cell(5, 3, "初始化伤害")
    ws.cell(5, 4, "设管线伤害(基础伤害)")
    ws.cell(5, 5, -1)
    ws.cell(5, 6, 0)
    steps = 加载战斗流程(wb)
    assert steps["伤害PVE"][0].成功跳转 == -1


def test_loader_rejects_bad_jump():
    wb = Workbook()
    ws = wb.active
    ws.title = "战斗流程"
    ws.cell(1, 1, "战斗流程")
    ws.cell(3, 1, "伤害PVE")
    for i, h in enumerate(("管线", "步骤序", "步骤名", "判定公式", "成功跳转", "失败跳转")):
        ws.cell(4, 1 + i, h)
    ws.cell(5, 1, "伤害PVE")
    ws.cell(5, 2, 1)
    ws.cell(5, 3, "初始化伤害")
    ws.cell(5, 4, "设管线伤害(基础伤害)")
    ws.cell(5, 5, 99)
    ws.cell(5, 6, 0)
    with pytest.raises(战斗流程加载错误):
        加载战斗流程(wb)


def test_stub_shield_then_hp():
    from 战斗模拟.管道 import 扣血 as hp_mod

    shields = [
        {
            "absorb_type": "全伤害",
            "remaining_capacity": 30.0,
            "remaining_hits": 0,
            "hits_only": False,
            "hit_capped": False,
        }
    ]
    out = hp_mod.解析扣血PVE(
        {
            "当前生命": 100.0,
            "最终伤害": 50.0,
            "伤害类型": "物理",
            "shields": shields,
        }
    )
    assert out["护盾吸收量"] == 30.0
    assert out["结算伤害"] == 20.0
    assert out["当前生命"] == 80.0
    assert out["已死亡"] is False
    assert out["本次破盾数"] == 1

    inv = hp_mod.解析扣血PVE(
        {"当前生命": 100.0, "最终伤害": 50.0, "无敌": True}
    )
    assert inv["结算伤害"] == 0.0
    assert inv["当前生命"] == 100.0
    assert inv["已死亡"] is False

    heal = hp_mod.解析回血PVE(
        {"当前生命": 80.0, "治疗值": 50.0, "生命值": 100.0}
    )
    assert heal["生命空档"] == 20.0
    assert heal["实际治疗"] == 20.0
    assert heal["过量治疗"] == 30.0
    assert heal["当前生命"] == 100.0
