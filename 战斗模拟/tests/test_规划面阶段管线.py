# -*- coding: utf-8 -*-
from __future__ import annotations

import random
from pathlib import Path

import pytest

from 战斗模拟.内核.流程解释器 import 流程解释器, 构建上下文
from 战斗模拟.load.战斗流程 import 加载事件接入规划, 加载战斗流程, 加载阶段规划, 是规划面

ROOT = Path(__file__).resolve().parents[3]
SANDBOX = ROOT / "数值框架" / "沙箱" / "战斗流程.xlsx"
OFFICIAL = ROOT / "数值框架" / "战斗数值框架.xlsx"


@pytest.fixture(scope="module")
def sandbox_path() -> Path:
    if not SANDBOX.is_file():
        pytest.skip(f"无规划面沙盒: {SANDBOX}")
    return SANDBOX


def _dmg_ctx():
    ctx = 构建上下文(
        基础伤害=100.0,
        伤害类型="物理",
        伤害来源="直接",
        仇恨系数=1.0,
        标记={"必中"},
        rng=random.Random(42),
        攻方={"伤害效果(PVE)": 1.0, "克制效果(PVE)": 1.0, "生命": 10000.0},
        守方={
            "生命": 50000.0,
            "生命上限": 50000.0,
            "闪避效果(PVE)": 0.0,
            "格挡效果(PVE)": 0.0,
            "暴击效果(PVE)": 0.0,
            "物理免伤效果(PVE)": 0.0,
        },
    )
    ctx["命中位"] = "必中"
    ctx["模式"] = "PVE"
    return ctx


def test_沙盒是规划面(sandbox_path: Path):
    from openpyxl import load_workbook

    wb = load_workbook(sandbox_path, data_only=True)
    try:
        assert 是规划面(wb["战斗流程"])
        planned = 加载阶段规划(wb)
        assert set(planned) >= {
            "伤害PVE",
            "伤害PVP",
            "治疗PVE",
            "治疗PVP",
            "状态PVE",
            "状态PVP",
            "特殊事件",
            "效果存续",
            "驱散与偷取",
            "施法",
        }
        assert 加载战斗流程(wb) == {}
        assert "机制流程" not in wb.sheetnames
        assert "事件" not in wb.sheetnames
        from 战斗模拟.load.战斗流程 import _规划面布局

        title_row, header_row, _ = _规划面布局(wb["战斗流程"])
        titles = [wb["战斗流程"].cell(title_row, c).value for c in range(1, (wb["战斗流程"].max_column or 1) + 1)]
        titles = [str(v).strip() for v in titles if v]
        assert titles[:4] == ["管线目录", "事件目录", "伤害PVE", "伤害PVP"]
        assert "特殊事件" in titles
        headers = [wb["战斗流程"].cell(header_row, c).value for c in range(1, (wb["战斗流程"].max_column or 1) + 1)]
        assert "启用" not in headers
        dmg_names = [s.段名 for s in planned["伤害PVE"]]
        heal_names = [s.段名 for s in planned["治疗PVE"]]
        assert "判定锁1血" not in dmg_names
        assert "判定结算目标免疫或无敌" not in dmg_names
        assert "判定禁疗" not in heal_names
        assert any("闪避效果(PVE)" in (s.判定公式 or "") for s in planned["伤害PVE"])
        assert any("闪避效果(PVP)" in (s.判定公式 or "") for s in planned["伤害PVP"])
    finally:
        wb.close()


def test_规划面伤害对齐旧食谱(sandbox_path: Path):
    from openpyxl import load_workbook

    if not OFFICIAL.is_file():
        pytest.skip(f"无正式簿: {OFFICIAL}")
    wb = load_workbook(OFFICIAL, data_only=True)
    try:
        if 是规划面(wb["战斗流程"]):
            old = 流程解释器(工作簿路径=OFFICIAL).执行管线("伤害PVE", _dmg_ctx())
        else:
            old_pipes = 加载战斗流程(wb)
            old_name = "伤害环境" if "伤害环境" in old_pipes else "伤害PVE"
            old = 流程解释器(管线表=old_pipes).执行管线(old_name, _dmg_ctx())
    finally:
        wb.close()
    new = 流程解释器(工作簿路径=sandbox_path).执行管线("伤害PVE", _dmg_ctx())
    assert new.管线 == "伤害PVE"
    assert new.最终伤害 == pytest.approx(old.最终伤害)
    assert new.实际扣血 == pytest.approx(old.实际扣血)
    names = [x.get("步骤名") for x in new.步骤轨迹]
    assert "判定锁1血" not in names
    assert names[-1] in {"写出伤害管线结果", "写出伤害结果"}


def test_结算入口与事件目录分流(sandbox_path: Path):
    from openpyxl import load_workbook

    wb = load_workbook(sandbox_path, data_only=True)
    try:
        ev = 加载事件接入规划(wb)
    finally:
        wb.close()
    assert ev["IMPACT"].进管线 == "伤害"
    assert ev["DOT_TICK"].进管线 == "伤害" and ev["DOT_TICK"].来源覆盖 == "持续"
    assert ev["AURA_APPLY"].进管线 == "状态施加"
    assert ev["CAST_START"].进管线 == "施法"
    assert ev["DISPEL"].进管线 == "驱散与偷取"


def test_施法完成可进入伤害(sandbox_path: Path):
    ctx = _dmg_ctx()
    ctx["事件种类"] = "CAST_COMPLETE"
    ctx.setdefault("标记", set()).add("请求伤害管线")
    out = 流程解释器(工作簿路径=sandbox_path).执行管线("施法", ctx)
    names = [x.get("步骤名") for x in out.步骤轨迹]
    assert "进入伤害结算" in names
    assert out.最终伤害 == pytest.approx(100.0)


def test_特殊事件不在基础管线(sandbox_path: Path):
    ctx = _dmg_ctx()
    ctx["守方"]["生命"] = 50.0
    ctx["守方"]["状态"] = {"锁1血"}
    ctx["当前生命"] = 50.0
    ctx["结算目标生命"] = 50.0
    out = 流程解释器(工作簿路径=sandbox_path).执行管线("伤害PVE", ctx)
    names = [x.get("步骤名") for x in out.步骤轨迹]
    assert "判定锁1血" in names
    assert out.上下文.get("当前生命") == pytest.approx(1.0)


def test_禁疗走特殊事件(sandbox_path: Path):
    ctx = 构建上下文(
        基础治疗=80.0,
        标记=set(),
        rng=random.Random(1),
        攻方={"施法治疗效果(PVE)": 1.0},
        守方={
            "生命": 100.0,
            "生命上限": 10000.0,
            "生命值": 10000.0,
            "受治疗效果(PVE)": 1.0,
            "状态": {"禁止受疗"},
        },
    )
    ctx["模式"] = "PVE"
    out = 流程解释器(工作簿路径=sandbox_path).执行管线("治疗环境", ctx)
    names = [x.get("步骤名") for x in out.步骤轨迹]
    assert "判定禁疗" in names
    from openpyxl import load_workbook

    wb = load_workbook(sandbox_path, data_only=True)
    try:
        heal_names = [s.段名 for s in 加载阶段规划(wb)["治疗PVE"]]
    finally:
        wb.close()
    assert "判定禁疗" not in heal_names
    assert out.实际治疗 == pytest.approx(0.0)
