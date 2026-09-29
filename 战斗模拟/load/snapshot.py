"""开场只取一次面板。战斗中的属性变化不再回表。"""
from __future__ import annotations

import openpyxl

from 战斗模拟.config import SHEET_MODEL

# 开场抄全部已经算出的终值，战斗中不再回表。


def _numeric(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return None


def read_cached_panels(path: str) -> dict[str, dict[str, float]]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[SHEET_MODEL]
        return _read_panel(ws)
    finally:
        wb.close()


def read_calculated_panels(path: str) -> dict[str, dict[str, float]]:
    """Excel 算完一次，只把终值抄进内存。不保存工作簿。"""
    import win32com.client

    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    wb = excel.Workbooks.Open(path)
    try:
        excel.CalculateFull()
        ws = wb.Worksheets(SHEET_MODEL)
        panels: dict[str, dict[str, float]] = {}
        builds = {}
        for col in range(29, 50):
            name = ws.Cells(3, col).Value
            if not name or str(name) in ("属性用途名", "标准·60级"):
                if builds:
                    break
                continue
            builds[col] = str(name)
        rows = {}
        for row in range(4, 220):
            label = ws.Cells(row, 29).Value
            if label:
                rows[str(label)] = row
        for col, build in builds.items():
            stats = {}
            for stat, row in rows.items():
                number = _numeric(ws.Cells(row, col).Value)
                if number is not None:
                    stats[stat] = number
            panels[build] = stats
        return panels
    finally:
        wb.Close(SaveChanges=False)
        excel.Quit()


def opening_panels(path: str, cached: dict[str, dict[str, float]]) -> dict[str, dict[str, float]]:
    if any(stats.get("生命值") for stats in cached.values()):
        return cached
    try:
        live = read_cached_panels(path)
    except Exception:
        live = {}
    if any(stats.get("生命值") for stats in live.values()):
        return live
    try:
        return read_calculated_panels(path)
    except Exception as exc:
        raise RuntimeError("开场面板需要 Excel 算一次终值，这次没有算出来") from exc


def _read_panel(ws) -> dict[str, dict[str, float]]:
    builds = {}
    for col in range(29, 50):
        name = ws.cell(3, col).value
        if not name or str(name) in ("属性用途名", "标准·60级"):
            if builds:
                break
            continue
        builds[col] = str(name)
    rows = {}
    for row in range(4, 220):
        label = ws.cell(row, 29).value
        if label:
            rows[str(label)] = row
    panels = {}
    for col, build in builds.items():
        stats = {}
        for stat, row in rows.items():
            number = _numeric(ws.cell(row, col).value)
            if number is not None:
                stats[stat] = number
        panels[build] = stats
    return panels
